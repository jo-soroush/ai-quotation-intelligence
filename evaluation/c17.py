"""V1-C17 deterministic, offline evaluation harness.

The harness grades existing public contracts against one fixed, synthetic
Golden Dataset.  It owns no production behavior and deliberately has no AWS,
network, Bedrock, or model-judge integration.
"""

from __future__ import annotations

import argparse
from collections.abc import Iterable, Mapping, Sequence
from contextlib import redirect_stdout
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
from typing import Any, Final

from openpyxl import load_workbook
from pydantic import ValidationError

from ai_quotation_intelligence.agent_tools import (
    AgentTools,
    RiskEvidenceInput,
    RiskEvidenceOutput,
)
from ai_quotation_intelligence.calculation import calculate_quote_total
from ai_quotation_intelligence.comparison import compare_historical_quote
from ai_quotation_intelligence.domain import (
    AgentRequest,
    AgentResult,
    AgentResultStatus,
    ApprovalDecision,
    ApprovalStatus,
    HistoricalQuote,
    Money,
    NewQuoteRequest,
    Quote,
)
from ai_quotation_intelligence.excel_export import SHEETS, export_approved_quote
from ai_quotation_intelligence.human_review import ReviewSession
from ai_quotation_intelligence.quotation_agent import QuotationAgent
from ai_quotation_intelligence.retrieval import retrieve_similar_quotes
from ai_quotation_intelligence.risk_evidence import build_risk_evidence


REPORT_SCHEMA_VERSION: Final = "1"
DATASET_SCHEMA_VERSION: Final = "1"
SUPPORTED_DATASET_VERSION: Final = "c17-golden-v1"
DEFAULT_DATASET_PATH: Final = Path(__file__).with_name("golden_v1.json")
PASS_CLAIM: Final = (
    "The system passed the defined deterministic contracts on the fixed synthetic Golden Dataset."
)
FAIL_CLAIM: Final = (
    "The system did not pass every defined deterministic contract on the fixed synthetic Golden Dataset."
)
ERROR_CLAIM: Final = "Evaluation integrity failed; no contract-pass claim is made."

GATING_THRESHOLDS: Final[dict[str, Decimal]] = {
    "calculation_correctness": Decimal("1"),
    "historical_comparison_correctness": Decimal("1"),
    "retrieval_hit_at_3": Decimal("1"),
    "evidence_provenance_correctness": Decimal("1"),
    "unsupported_risk_rate": Decimal("0"),
    "structured_output_validity": Decimal("1"),
    "tool_call_contract_success": Decimal("1"),
    "agent_task_completion": Decimal("1"),
    "excel_correctness": Decimal("1"),
}
REPORT_ONLY_METRICS: Final = ("latency", "bedrock_usage", "cost")
CASE_CATEGORIES: Final = (
    "calculation",
    "historical_comparison",
    "retrieval",
    "risk_evidence",
    "structured_output",
    "tool_call",
    "agent_task",
    "excel",
)
_CATEGORY_METRICS: Final = {
    "calculation": "calculation_correctness",
    "historical_comparison": "historical_comparison_correctness",
    "retrieval": "retrieval_hit_at_3",
    "risk_evidence": "evidence_provenance_correctness",
    "structured_output": "structured_output_validity",
    "tool_call": "tool_call_contract_success",
    "agent_task": "agent_task_completion",
    "excel": "excel_correctness",
}
_PROHIBITED_REPORT_KEYS: Final = frozenset({
    "aws_access_key_id",
    "aws_secret_access_key",
    "aws_session_token",
    "secret",
    "token",
    "raw_prompt",
    "prompt",
    "raw_model_response",
    "model_response",
    "reviewer_id",
    "workbook",
    "workbook_bytes",
    "request_body",
    "response_body",
    "customer_data",
})
_BOUNDED_FAILURE_REASONS: Final = frozenset({
    "calculation_mismatch",
    "historical_comparison_mismatch",
    "expected_id_not_in_top_3",
    "risk_evidence_or_provenance_mismatch",
    "tool_contract_mismatch",
    "scripted_agent_contract_mismatch",
    "structured_output_invalid",
    "workbook_contract_mismatch",
})


class EvaluationIntegrityError(ValueError):
    """The dataset or report cannot support an interpretable evaluation."""


@dataclass(frozen=True)
class LoadedDataset:
    data: dict[str, Any]
    sha256: str
    path: Path


@dataclass(frozen=True)
class _ScriptedResponse:
    status: AgentResultStatus
    text: str
    request_id: str


class _ScriptedModel:
    """Provider-neutral fixed trace; it has no external-call capability."""

    def __init__(self, script: Sequence[Mapping[str, Any]]) -> None:
        self._outputs = [
            json.dumps(item, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
            for item in script
        ]

    def invoke(self, _prompt: str, *, request_id: str) -> _ScriptedResponse:
        if not self._outputs:
            raise RuntimeError("script_exhausted")
        return _ScriptedResponse(
            status=AgentResultStatus.SUCCESS,
            text=self._outputs.pop(0),
            request_id=request_id,
        )


def _require(condition: bool, reason: str) -> None:
    if not condition:
        raise EvaluationIntegrityError(reason)


def _decimal(value: Any, reason: str) -> Decimal:
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        raise EvaluationIntegrityError(reason) from None
    if not result.is_finite():
        raise EvaluationIntegrityError(reason)
    return result


def _optional_decimal(value: Any, reason: str) -> Decimal | None:
    return None if value is None else _decimal(value, reason)


def _sha_path(path: Path) -> Path:
    return path.with_suffix(".sha256")


def _validate_dataset_shape(data: Any) -> dict[str, Any]:
    _require(isinstance(data, dict), "dataset_not_object")
    _require(set(data) == {
        "dataset_schema_version", "dataset_version", "synthetic", "oracle_source",
        "history", "cases",
    }, "dataset_fields_invalid")
    _require(data["dataset_schema_version"] == DATASET_SCHEMA_VERSION,
             "unsupported_dataset_schema_version")
    _require(data["dataset_version"] == SUPPORTED_DATASET_VERSION,
             "unsupported_dataset_version")
    _require(data["synthetic"] is True, "dataset_not_synthetic")
    _require(data["oracle_source"] == "independently_authored_and_preverified",
             "oracle_source_invalid")
    _require(isinstance(data["history"], list) and data["history"], "history_invalid")
    _require(isinstance(data["cases"], dict) and set(data["cases"]) == set(CASE_CATEGORIES),
             "case_categories_invalid")

    history_ids: list[str] = []
    source_ids: list[str] = []
    for raw in data["history"]:
        try:
            record = HistoricalQuote.model_validate(raw)
        except (ValidationError, ValueError, TypeError):
            raise EvaluationIntegrityError("historical_fixture_invalid") from None
        _require(record.data_origin.value == "synthetic", "history_not_synthetic")
        history_ids.append(record.quote.quote_id)
        source_ids.append(record.source_id)
    _require(len(history_ids) == len(set(history_ids)), "duplicate_history_quote_id")
    _require(len(source_ids) == len(set(source_ids)), "duplicate_history_source_id")

    all_case_ids: list[str] = []
    for category in CASE_CATEGORIES:
        cases = data["cases"][category]
        _require(isinstance(cases, list) and cases, f"{category}_cases_invalid")
        for case in cases:
            _require(isinstance(case, dict), "case_not_object")
            case_id = case.get("case_id")
            _require(isinstance(case_id, str) and case_id.strip() == case_id and case_id,
                     "case_id_invalid")
            all_case_ids.append(case_id)
    _require(len(all_case_ids) == len(set(all_case_ids)), "duplicate_case_id")

    _validate_case_contracts(data, set(history_ids))
    return data


def _validate_case_contracts(data: dict[str, Any], history_ids: set[str]) -> None:
    cases = data["cases"]
    for case in cases["calculation"]:
        _require(set(case) == {"case_id", "quote", "expected_total"},
                 "calculation_case_fields_invalid")
        try:
            quote = Quote.model_validate(case["quote"])
            Money.model_validate(case["expected_total"])
        except (ValidationError, ValueError, TypeError):
            raise EvaluationIntegrityError("calculation_case_invalid") from None
        _require(quote.data_origin.value == "synthetic", "calculation_case_not_synthetic")

    for case in cases["historical_comparison"]:
        _require(set(case) == {"case_id", "history_quote_id", "expected"},
                 "comparison_case_fields_invalid")
        _require(case["history_quote_id"] in history_ids, "comparison_history_missing")
        expected = case["expected"]
        _require(isinstance(expected, dict) and set(expected) == {
            "quote_id", "source_id", "hour_variance", "cost_variance",
            "scope_changed", "outcome_label",
        }, "comparison_expected_invalid")
        for metric in ("hour_variance", "cost_variance"):
            value = expected[metric]
            _require(isinstance(value, dict) and set(value) == {
                "estimated", "actual", "variance", "percentage", "unit",
            }, "comparison_variance_expected_invalid")
            for field in ("estimated", "actual", "variance", "percentage"):
                _decimal(value[field], "comparison_number_invalid")

    for case in cases["retrieval"]:
        _require(set(case) == {
            "case_id", "request", "history_quote_ids", "expected_similar_quote_ids",
        }, "retrieval_case_fields_invalid")
        try:
            NewQuoteRequest.model_validate(case["request"])
        except (ValidationError, ValueError, TypeError):
            raise EvaluationIntegrityError("retrieval_request_invalid") from None
        _validate_history_references(case["history_quote_ids"], history_ids)
        expected = case["expected_similar_quote_ids"]
        _require(isinstance(expected, list) and expected and all(
            isinstance(item, str) and item in history_ids for item in expected
        ), "retrieval_expected_ids_invalid")

    for case in cases["risk_evidence"]:
        _require(set(case) == {
            "case_id", "metric", "history_quote_ids", "required_evidence", "expected",
        }, "risk_case_fields_invalid")
        _require(case["metric"] in {"hours", "cost"}, "risk_metric_invalid")
        _require(type(case["required_evidence"]) is bool, "required_evidence_invalid")
        _validate_history_references(case["history_quote_ids"], history_ids)
        expected = case["expected"]
        _require(isinstance(expected, dict) and set(expected) == {
            "status", "comparable_project_count", "overrun_count", "overrun_rate",
            "average_variance", "median_variance", "source_quote_ids", "source_ids",
            "evidence_ids",
        }, "risk_expected_invalid")
        for field in ("overrun_rate", "average_variance", "median_variance"):
            _optional_decimal(expected[field], "risk_expected_number_invalid")
        _require(isinstance(expected["evidence_ids"], list), "risk_expected_evidence_invalid")

    for case in cases["tool_call"]:
        _require(set(case) == {
            "case_id", "tool", "arguments", "expected_status", "expected_evidence_id",
        }, "tool_case_fields_invalid")
        _require(case["tool"] == "get_risk_evidence", "tool_case_unsupported")
        _require(case["arguments"] in ({"metric": "hours"}, {"metric": "cost"}),
                 "tool_arguments_invalid")

    agent_ids: set[str] = set()
    for case in cases["agent_task"]:
        _require(set(case) == {
            "case_id", "request_id", "quotation_request", "script", "expected",
        }, "agent_case_fields_invalid")
        try:
            NewQuoteRequest.model_validate(case["quotation_request"])
        except (ValidationError, ValueError, TypeError):
            raise EvaluationIntegrityError("agent_request_invalid") from None
        _require(isinstance(case["script"], list) and case["script"], "agent_script_invalid")
        _require(isinstance(case["expected"], dict) and set(case["expected"]) == {
            "status", "quote_id", "estimated_total", "evidence_ids",
        }, "agent_expected_invalid")
        _decimal(case["expected"]["estimated_total"], "agent_expected_total_invalid")
        agent_ids.add(case["case_id"])

    for category, exact_fields in (
        ("structured_output", {"case_id", "agent_case_id", "expected_schema", "expected_status"}),
        ("excel", {"case_id", "agent_case_id", "expected_sheets", "expected_quote_id", "expected_total"}),
    ):
        for case in cases[category]:
            _require(set(case) == exact_fields, f"{category}_case_fields_invalid")
            _require(case["agent_case_id"] in agent_ids, f"{category}_agent_case_missing")
            if category == "excel":
                _decimal(case["expected_total"], "excel_expected_total_invalid")


def _validate_history_references(values: Any, known: set[str]) -> None:
    _require(isinstance(values, list) and values, "history_references_invalid")
    _require(len(values) == len(set(values)) and all(
        isinstance(item, str) and item in known for item in values
    ), "history_references_invalid")


def load_dataset(path: Path | str = DEFAULT_DATASET_PATH) -> LoadedDataset:
    """Load and authenticate a Golden Dataset using its sibling SHA-256 file."""

    dataset_path = Path(path)
    try:
        payload = dataset_path.read_bytes()
        expected = _sha_path(dataset_path).read_text(encoding="ascii").strip()
    except OSError:
        raise EvaluationIntegrityError("dataset_or_hash_unreadable") from None
    digest = sha256(payload).hexdigest()
    _require(len(expected) == 64 and all(char in "0123456789abcdef" for char in expected),
             "dataset_hash_record_invalid")
    _require(digest == expected, "dataset_hash_mismatch")
    try:
        parsed = json.loads(payload)
    except (json.JSONDecodeError, UnicodeDecodeError):
        raise EvaluationIntegrityError("dataset_json_invalid") from None
    return LoadedDataset(_validate_dataset_shape(parsed), digest, dataset_path)


def _history(data: dict[str, Any]) -> dict[str, HistoricalQuote]:
    return {
        raw["quote"]["quote_id"]: HistoricalQuote.model_validate(raw)
        for raw in data["history"]
    }


def _case_result(
    case_id: str,
    metric: str,
    passed: bool,
    reason: str,
    *,
    evidence_refs: Iterable[str] = (),
) -> dict[str, Any]:
    return {
        "case_id": case_id,
        "metric": metric,
        "status": "PASS" if passed else "FAIL",
        "failure_reason": None if passed else reason,
        "evidence_references": sorted(set(evidence_refs)),
    }


def _variance_matches(actual: Any, expected: Mapping[str, Any]) -> bool:
    return (
        actual is not None
        and actual.estimated == _decimal(expected["estimated"], "comparison_number_invalid")
        and actual.actual == _decimal(expected["actual"], "comparison_number_invalid")
        and actual.variance == _decimal(expected["variance"], "comparison_number_invalid")
        and actual.percentage == _decimal(expected["percentage"], "comparison_number_invalid")
        and actual.unit == expected["unit"]
    )


def _grade_calculation(cases: list[dict[str, Any]]) -> list[dict[str, Any]]:
    results = []
    for case in cases:
        quote = Quote.model_validate(case["quote"])
        expected = Money.model_validate(case["expected_total"])
        passed = calculate_quote_total(quote) == expected
        results.append(_case_result(
            case["case_id"], "calculation_correctness", passed, "calculation_mismatch",
        ))
    return results


def _grade_comparison(
    cases: list[dict[str, Any]], history: Mapping[str, HistoricalQuote]
) -> list[dict[str, Any]]:
    results = []
    for case in cases:
        actual = compare_historical_quote(history[case["history_quote_id"]])
        expected = case["expected"]
        passed = (
            actual.quote_id == expected["quote_id"]
            and actual.source_id == expected["source_id"]
            and _variance_matches(actual.hour_variance, expected["hour_variance"])
            and _variance_matches(actual.cost_variance, expected["cost_variance"])
            and actual.scope_changed == expected["scope_changed"]
            and actual.outcome_label == expected["outcome_label"]
        )
        results.append(_case_result(
            case["case_id"], "historical_comparison_correctness", passed,
            "historical_comparison_mismatch",
            evidence_refs=(actual.quote_id, actual.source_id),
        ))
    return results


def _grade_retrieval(
    cases: list[dict[str, Any]], history: Mapping[str, HistoricalQuote]
) -> list[dict[str, Any]]:
    results = []
    for case in cases:
        request = NewQuoteRequest.model_validate(case["request"])
        records = tuple(history[item] for item in case["history_quote_ids"])
        returned = retrieve_similar_quotes(request, records, limit=3)
        top_three = [item.result.quote_id for item in returned[:3]]
        expected = set(case["expected_similar_quote_ids"])
        passed = bool(expected & set(top_three))
        results.append(_case_result(
            case["case_id"], "retrieval_hit_at_3", passed, "expected_id_not_in_top_3",
            evidence_refs=top_three,
        ))
    return results


def _grade_risk(
    cases: list[dict[str, Any]], history: Mapping[str, HistoricalQuote]
) -> tuple[list[dict[str, Any]], int, int]:
    results: list[dict[str, Any]] = []
    unsupported = 0
    total_references = 0
    for case in cases:
        records = tuple(history[item] for item in case["history_quote_ids"])
        report = build_risk_evidence(records, metric=case["metric"])
        expected = case["expected"]
        evidence_ids = [item.evidence_id for item in report.evidence]
        expected_ids = set(expected["evidence_ids"])
        expected_quotes = set(expected["source_quote_ids"])
        total_references += len(report.evidence)
        unsupported += sum(
            item.evidence_id not in expected_ids
            or not item.source_quote_ids
            or not set(item.source_quote_ids) <= expected_quotes
            for item in report.evidence
        )
        expected_match = (
            report.status.value == expected["status"]
            and report.comparable_project_count == expected["comparable_project_count"]
            and report.overrun_count == expected["overrun_count"]
            and report.overrun_rate == _optional_decimal(expected["overrun_rate"], "risk_number_invalid")
            and report.average_variance == _optional_decimal(expected["average_variance"], "risk_number_invalid")
            and report.median_variance == _optional_decimal(expected["median_variance"], "risk_number_invalid")
            and list(report.source_quote_ids) == expected["source_quote_ids"]
            and list(report.source_ids) == expected["source_ids"]
            and evidence_ids == expected["evidence_ids"]
            and all(item.data_origin.value == "synthetic" for item in report.evidence)
        )
        required_present = not case["required_evidence"] or bool(report.evidence)
        passed = expected_match and required_present
        results.append(_case_result(
            case["case_id"], "evidence_provenance_correctness", passed,
            "risk_evidence_or_provenance_mismatch",
            evidence_refs=(*report.source_quote_ids, *report.source_ids, *evidence_ids),
        ))
    return results, unsupported, total_references


def _grade_tools(cases: list[dict[str, Any]]) -> list[dict[str, Any]]:
    results = []
    tools = AgentTools()
    for case in cases:
        request_id = f"c17-{case['case_id']}"
        request = RiskEvidenceInput(request_id=request_id, **case["arguments"])
        output = tools.resolve_tool(case["tool"])(request)
        passed = (
            type(output) is RiskEvidenceOutput
            and output.request_id == request_id
            and output.tool == case["tool"]
            and output.report.status.value == case["expected_status"]
            and case["expected_evidence_id"] in {
                item.evidence_id for item in output.report.evidence
            }
            and all(item.source_quote_ids for item in output.report.evidence)
        )
        results.append(_case_result(
            case["case_id"], "tool_call_contract_success", passed,
            "tool_contract_mismatch",
            evidence_refs=(case["expected_evidence_id"],),
        ))
    return results


def _run_agent_cases(
    cases: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, AgentResult]]:
    results = []
    outputs: dict[str, AgentResult] = {}
    for case in cases:
        request = AgentRequest(
            request_id=case["request_id"],
            quotation_request=NewQuoteRequest.model_validate(case["quotation_request"]),
        )
        result = QuotationAgent(_ScriptedModel(case["script"]), AgentTools()).run(request)
        outputs[case["case_id"]] = result
        expected = case["expected"]
        quote = result.draft_quote.quote if result.draft_quote is not None else None
        passed = (
            result.status.value == expected["status"]
            and quote is not None
            and quote.quote_id == expected["quote_id"]
            and quote.estimated_total_cost is not None
            and quote.estimated_total_cost.amount
            == _decimal(expected["estimated_total"], "agent_expected_total_invalid")
            and result.evidence_ids == expected["evidence_ids"]
            and result.draft_quote.evidence_ids == expected["evidence_ids"]
            and all(set(item.evidence_ids) <= set(result.evidence_ids)
                    for item in result.risk_suggestions)
        )
        results.append(_case_result(
            case["case_id"], "agent_task_completion", passed,
            "scripted_agent_contract_mismatch",
            evidence_refs=result.evidence_ids,
        ))
    return results, outputs


def _grade_structured(
    cases: list[dict[str, Any]], agent_results: Mapping[str, AgentResult]
) -> list[dict[str, Any]]:
    results = []
    for case in cases:
        source = agent_results[case["agent_case_id"]]
        try:
            clean = AgentResult.model_validate(source.model_dump())
            passed = (
                case["expected_schema"] == "AgentResult"
                and clean.status.value == case["expected_status"]
                and clean.request_id == source.request_id
            )
        except (ValidationError, ValueError, TypeError):
            passed = False
        results.append(_case_result(
            case["case_id"], "structured_output_validity", passed,
            "structured_output_invalid",
            evidence_refs=source.evidence_ids,
        ))
    return results


def _grade_excel(
    cases: list[dict[str, Any]], agent_results: Mapping[str, AgentResult]
) -> list[dict[str, Any]]:
    results = []
    for case in cases:
        source = agent_results[case["agent_case_id"]]
        if source.draft_quote is None:
            results.append(_case_result(
                case["case_id"], "excel_correctness", False,
                "workbook_contract_mismatch",
            ))
            continue
        session = ReviewSession(source)
        session.decide(
            ApprovalDecision(
                quote_id=source.draft_quote.quote.quote_id,
                status=ApprovalStatus.APPROVED,
                reviewer_id="c17-synthetic-reviewer",
                decided_at=datetime(2026, 2, 3, 10, 0, tzinfo=timezone.utc),
            ),
            current=source,
        )
        blob = export_approved_quote(session, current=source)
        book = load_workbook(BytesIO(blob), data_only=False, keep_links=True)
        quote_sheet = book["Quotation"] if "Quotation" in book.sheetnames else None
        passed = (
            book.sheetnames == case["expected_sheets"] == list(SHEETS)
            and quote_sheet is not None
            and quote_sheet["B2"].value == case["expected_quote_id"]
            and Decimal(str(quote_sheet.cell(quote_sheet.max_row, 5).value))
            == _decimal(case["expected_total"], "excel_expected_total_invalid")
            and not book._external_links
            and book.vba_archive is None
        )
        book.close()
        results.append(_case_result(
            case["case_id"], "excel_correctness", passed,
            "workbook_contract_mismatch",
            evidence_refs=(case["expected_quote_id"],),
        ))
    return results


def _gating_metric(
    name: str,
    case_results: Sequence[dict[str, Any]],
    expected_count: int,
    *,
    numerator: int | None = None,
    denominator: int | None = None,
) -> dict[str, Any]:
    _require(name in GATING_THRESHOLDS, "unknown_metric_definition")
    _require(len(case_results) == expected_count, "case_count_mismatch")
    _require(len({item["case_id"] for item in case_results}) == expected_count,
             "case_result_identity_mismatch")
    provided_counts = numerator is not None or denominator is not None
    _require(
        name == "unsupported_risk_rate" or not provided_counts,
        "metric_denominator_manipulation",
    )
    if numerator is None:
        numerator = sum(item["status"] == "PASS" for item in case_results)
    if denominator is None:
        denominator = expected_count
    _require(type(numerator) is int and type(denominator) is int,
             "metric_count_type_invalid")
    _require(0 <= numerator <= denominator, "metric_count_invalid")
    _require(denominator > 0 or name == "unsupported_risk_rate",
             "metric_denominator_invalid")
    value = Decimal("0") if denominator == 0 else Decimal(numerator) / Decimal(denominator)
    threshold = GATING_THRESHOLDS[name]
    passed = value <= threshold if name == "unsupported_risk_rate" else value >= threshold
    return {
        "name": name,
        "report_only": False,
        "numerator": numerator,
        "denominator": denominator,
        "value": str(value),
        "threshold": str(threshold),
        "status": "PASS" if passed else "FAIL",
    }


def _report_only_metric(name: str) -> dict[str, Any]:
    _require(name in REPORT_ONLY_METRICS, "unknown_report_only_metric")
    return {
        "name": name,
        "report_only": True,
        "numerator": None,
        "denominator": None,
        "value": None,
        "threshold": None,
        "status": "NOT_MEASURED_OFFLINE",
    }


def _validate_report(report: Mapping[str, Any]) -> None:
    _require(set(report) == {
        "report_schema_version", "dataset_version", "dataset_sha256", "synthetic",
        "overall_status", "error_category", "claim_boundary", "metrics", "case_results",
    }, "report_fields_invalid")
    _require(report["report_schema_version"] == REPORT_SCHEMA_VERSION,
             "report_schema_version_invalid")
    _require(report["dataset_version"] == SUPPORTED_DATASET_VERSION,
             "report_dataset_version_invalid")
    digest = report["dataset_sha256"]
    _require(isinstance(digest, str) and len(digest) == 64
             and all(char in "0123456789abcdef" for char in digest),
             "report_dataset_hash_invalid")
    _require(report["synthetic"] is True, "report_synthetic_marker_invalid")
    _require(report["overall_status"] in {"PASS", "FAIL", "ERROR"},
             "report_status_invalid")

    metrics = report["metrics"]
    _require(isinstance(metrics, list), "report_metrics_invalid")
    names = [item.get("name") for item in metrics if isinstance(item, dict)]
    _require(set(names) == set(GATING_THRESHOLDS) | set(REPORT_ONLY_METRICS)
             and len(names) == len(set(names)), "report_metric_set_invalid")
    for metric in metrics:
        _require(isinstance(metric, dict) and set(metric) == {
            "name", "report_only", "numerator", "denominator", "value", "threshold", "status",
        }, "report_metric_fields_invalid")
        if metric["name"] in GATING_THRESHOLDS:
            _require(metric["report_only"] is False, "gating_metric_marked_report_only")
            _require(metric["threshold"] == str(GATING_THRESHOLDS[metric["name"]]),
                     "gating_threshold_invalid")
            _require(metric["status"] in {"PASS", "FAIL"}, "gating_status_invalid")
        else:
            _require(metric["report_only"] is True, "report_only_metric_gating")
            _require(metric["threshold"] is None and metric["status"] == "NOT_MEASURED_OFFLINE",
                     "report_only_metric_invalid")

    case_results = report["case_results"]
    _require(isinstance(case_results, list), "report_case_results_invalid")
    for result in case_results:
        _require(isinstance(result, dict) and set(result) == {
            "case_id", "metric", "status", "failure_reason", "evidence_references",
        }, "report_case_fields_invalid")
        _require(result["status"] in {"PASS", "FAIL"}, "report_case_status_invalid")
        _require(result["failure_reason"] is None
                 or result["failure_reason"] in _BOUNDED_FAILURE_REASONS,
                 "report_failure_reason_unbounded")
        _require(isinstance(result["evidence_references"], list),
                 "report_evidence_references_invalid")
    gating = [item for item in metrics if not item["report_only"]]
    if report["overall_status"] == "PASS":
        _require(all(item["status"] == "PASS" for item in gating),
                 "misleading_pass_status")
        _require(all(item["status"] == "PASS" for item in case_results),
                 "misleading_pass_case_status")
    elif report["overall_status"] == "FAIL":
        _require(any(item["status"] == "FAIL" for item in gating),
                 "misleading_fail_status")
    expected_claim = {
        "PASS": PASS_CLAIM,
        "FAIL": FAIL_CLAIM,
        "ERROR": ERROR_CLAIM,
    }[report["overall_status"]]
    _require(report["claim_boundary"] == expected_claim, "claim_boundary_invalid")
    if report["overall_status"] == "ERROR":
        _require(report["error_category"] in {
            "evaluation_integrity_error", "harness_execution_error",
        }, "report_error_category_invalid")
    else:
        _require(report["error_category"] is None, "report_error_category_invalid")
    _validate_report_privacy(report)


def _validate_report_privacy(value: Any) -> None:
    if isinstance(value, Mapping):
        for key, child in value.items():
            _require(str(key).lower() not in _PROHIBITED_REPORT_KEYS,
                     "prohibited_report_field")
            _validate_report_privacy(child)
    elif isinstance(value, list):
        for child in value:
            _validate_report_privacy(child)
    elif isinstance(value, (bytes, bytearray, memoryview)):
        raise EvaluationIntegrityError("binary_report_value_forbidden")


def _error_report(dataset_hash: str, category: str) -> dict[str, Any]:
    return {
        "report_schema_version": REPORT_SCHEMA_VERSION,
        "dataset_version": SUPPORTED_DATASET_VERSION,
        "dataset_sha256": dataset_hash if len(dataset_hash) == 64 else "0" * 64,
        "synthetic": True,
        "overall_status": "ERROR",
        "error_category": category,
        "claim_boundary": ERROR_CLAIM,
        "metrics": [
            {
                "name": name,
                "report_only": False,
                "numerator": 0,
                "denominator": 0,
                "value": None,
                "threshold": str(threshold),
                "status": "FAIL",
            }
            for name, threshold in GATING_THRESHOLDS.items()
        ] + [_report_only_metric(name) for name in REPORT_ONLY_METRICS],
        "case_results": [],
    }


def evaluate(path: Path | str = DEFAULT_DATASET_PATH) -> dict[str, Any]:
    """Run the fixed suite and return a deterministic, privacy-bounded report."""

    path = Path(path)
    try:
        loaded = load_dataset(path)
        data = loaded.data
        history = _history(data)
        cases = data["cases"]

        calculation = _grade_calculation(cases["calculation"])
        comparisons = _grade_comparison(cases["historical_comparison"], history)
        retrieval = _grade_retrieval(cases["retrieval"], history)
        risk, unsupported, risk_references = _grade_risk(cases["risk_evidence"], history)
        tools = _grade_tools(cases["tool_call"])
        agent, agent_results = _run_agent_cases(cases["agent_task"])
        structured = _grade_structured(cases["structured_output"], agent_results)
        excel = _grade_excel(cases["excel"], agent_results)

        by_category = {
            "calculation": calculation,
            "historical_comparison": comparisons,
            "retrieval": retrieval,
            "risk_evidence": risk,
            "structured_output": structured,
            "tool_call": tools,
            "agent_task": agent,
            "excel": excel,
        }
        case_results = [
            result
            for category in CASE_CATEGORIES
            for result in by_category[category]
        ]
        metrics = [
            _gating_metric(
                _CATEGORY_METRICS[category],
                by_category[category],
                len(cases[category]),
            )
            for category in CASE_CATEGORIES
        ]
        metrics.append(_gating_metric(
            "unsupported_risk_rate",
            risk,
            len(cases["risk_evidence"]),
            numerator=unsupported,
            denominator=risk_references,
        ))
        metrics.extend(_report_only_metric(name) for name in REPORT_ONLY_METRICS)
        metric_names = [item["name"] for item in metrics]
        _require(set(metric_names) == set(GATING_THRESHOLDS) | set(REPORT_ONLY_METRICS)
                 and len(metric_names) == len(set(metric_names)),
                 "metric_coverage_invalid")
        overall_status = "PASS" if all(
            item["status"] == "PASS" for item in metrics if not item["report_only"]
        ) else "FAIL"
        report = {
            "report_schema_version": REPORT_SCHEMA_VERSION,
            "dataset_version": data["dataset_version"],
            "dataset_sha256": loaded.sha256,
            "synthetic": True,
            "overall_status": overall_status,
            "error_category": None,
            "claim_boundary": PASS_CLAIM if overall_status == "PASS" else FAIL_CLAIM,
            "metrics": sorted(metrics, key=lambda item: item["name"]),
            "case_results": sorted(case_results, key=lambda item: item["case_id"]),
        }
        _validate_report(report)
        return report
    except EvaluationIntegrityError:
        try:
            digest = sha256(path.read_bytes()).hexdigest()
        except OSError:
            digest = "0" * 64
        report = _error_report(digest, "evaluation_integrity_error")
        _validate_report(report)
        return report
    except Exception:
        try:
            digest = sha256(path.read_bytes()).hexdigest()
        except OSError:
            digest = "0" * 64
        report = _error_report(digest, "harness_execution_error")
        _validate_report(report)
        return report


def serialize_report(report: Mapping[str, Any]) -> str:
    """Return the authoritative stable JSON representation."""

    _validate_report(report)
    return json.dumps(report, ensure_ascii=False, sort_keys=True, indent=2) + "\n"


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run the offline V1-C17 evaluation")
    parser.add_argument("--dataset", type=Path, default=DEFAULT_DATASET_PATH)
    args = parser.parse_args(argv)
    # Existing production instrumentation writes operational JSON to stdout.
    # Keep the CLI's authoritative report a single JSON document without
    # changing or disabling production logging behavior.
    with redirect_stdout(StringIO()):
        report = evaluate(args.dataset)
    print(serialize_report(report), end="")
    return 0 if report["overall_status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())

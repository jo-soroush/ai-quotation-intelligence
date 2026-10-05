"""V1-C19 local, deterministic, synthetic Golden Case harness.

This non-production owner composes already-delivered public contracts. It has
no live-service client, persistence authority, business logic, or deployment
role. Generated reports and workbooks are ephemeral.
"""

from __future__ import annotations

import argparse
from collections import Counter
from collections.abc import Callable, Mapping, Sequence
from contextlib import redirect_stdout
from dataclasses import dataclass
from decimal import Decimal
from hashlib import sha256
from io import BytesIO, StringIO
import json
import logging
from pathlib import Path
from typing import Any, Final

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from evaluation import c17
from ai_quotation_intelligence.agent_tools import (
    AgentTools,
    ComparisonsOutput,
    HistoricalQuotesOutput,
    RiskEvidenceOutput,
    SimilarQuotesOutput,
    StatisticsOutput,
)
from ai_quotation_intelligence.api import LocalQuoteStore, create_app
from ai_quotation_intelligence.bedrock import BedrockConverseAdapter
from ai_quotation_intelligence.config import Settings
from ai_quotation_intelligence.domain import AgentResultStatus
from ai_quotation_intelligence.logging_config import configure_logging
from ai_quotation_intelligence.quotation_agent import QuotationAgent
from ai_quotation_intelligence.s3_storage import S3StorageAdapter, derive_object_key
from ai_quotation_intelligence.storage_contract import WorkbookArtifact


REPORT_SCHEMA_VERSION: Final = "c19-golden-report-v1"
SCENARIO_SCHEMA_VERSION: Final = "1"
SCENARIO_VERSION: Final = "c19-golden-v1"
DEFAULT_SCENARIO_PATH: Final = Path(__file__).with_name("golden_case_v1.json")
C17_DATASET_SHA256: Final = "1e2bec2a2c58a9c0082491d8ef746da723689a7f73a9c255e47268a6a6a4cc6b"
C17_REPORT_SHA256: Final = "928880c6042e4ab9fa826cf1ba51b88cd2a762209c5006144c2e53f233cf832c"
PASS_CLAIM: Final = (
    "One fixed synthetic end-to-end Golden Case passed across the existing V1 component contracts."
)
REQUIRED_STEPS: Final = (
    "fixture_identity",
    "fastapi_request_boundary",
    "core_calculation",
    "historical_retrieval",
    "similar_retrieval",
    "historical_comparison",
    "deterministic_statistics",
    "risk_evidence",
    "agent_tools",
    "provider_adapter",
    "quotation_agent",
    "structured_output_validation",
    "human_approval",
    "excel_generation",
    "excel_reload_reconciliation",
    "storage_adapter",
    "observability",
    "c17_regression",
)
TOOL_SEQUENCE: Final = (
    "get_historical_quotes",
    "find_similar_quotes",
    "compare_estimate_to_actual",
    "calculate_quote_statistics",
    "get_risk_evidence",
)
FAULTS: Final = frozenset({
    "provider_invalid",
    "provider_unavailable",
    "skip_approval",
    "excel_reconciliation_failure",
    "storage_false_success",
    "omit_observability",
})
_REPORT_KEYS: Final = frozenset({
    "report_schema_version", "scenario_id", "scenario_version", "scenario_sha256",
    "synthetic", "overall_status", "required_step_results",
    "deterministic_commercial_reconciliation", "retrieval_comparison_references",
    "risk_evidence_references", "agent_tool_trace_outcome",
    "structured_output_validity", "approval_workflow_result",
    "excel_reconciliation_result", "storage_contract_result",
    "observability_contract_result", "c17_regression_result",
    "bounded_failure_category", "claim_boundary", "limitations",
})
_PROHIBITED_KEYS: Final = frozenset({
    "aws_access_key_id", "aws_secret_access_key", "aws_session_token", "secret",
    "token", "prompt", "raw_prompt", "model_response", "raw_model_response",
    "reviewer_id", "workbook", "workbook_bytes", "request_body", "response_body",
    "raw_logs", "provider_exception", "customer_data",
})


class ScenarioIntegrityError(ValueError):
    """The fixed scenario cannot support a valid Golden Case evaluation."""


class ReportIntegrityError(ValueError):
    """The machine-readable evidence report violates its fixed contract."""


@dataclass(frozen=True)
class LoadedScenario:
    data: dict[str, Any]
    sha256: str
    path: Path


@dataclass(frozen=True)
class StepResult:
    step_id: str
    passed: bool
    failure_category: str | None = None

    def as_report(self) -> dict[str, Any]:
        return {
            "step_id": self.step_id,
            "status": "PASS" if self.passed else "FAIL",
            "passed": self.passed,
            "failure_category": self.failure_category,
        }


class _EventCapture(logging.Handler):
    def __init__(self) -> None:
        super().__init__()
        self.events: list[dict[str, Any]] = []

    def emit(self, record: logging.LogRecord) -> None:
        event = getattr(record, "aqi_event", None)
        if type(event) is dict:
            self.events.append(dict(event))


class _ScriptedConverseClient:
    """Fixed local provider client; it cannot reach a network."""

    def __init__(self, trace: Sequence[Mapping[str, Any]], fault: str | None = None) -> None:
        self._trace = [dict(item) for item in trace]
        self._fault = fault
        self.call_count = 0
        self.request_shapes_valid = True

    def converse(self, **kwargs: Any) -> dict[str, Any]:
        self.call_count += 1
        self.request_shapes_valid = self.request_shapes_valid and (
            set(kwargs) == {"modelId", "messages", "inferenceConfig"}
            and isinstance(kwargs.get("messages"), list)
        )
        if self._fault == "provider_unavailable":
            raise RuntimeError("synthetic provider unavailable")
        if self._fault == "provider_invalid":
            return {"output": {"message": {"content": []}}}
        if not self._trace:
            raise RuntimeError("script exhausted")
        text = json.dumps(self._trace.pop(0), sort_keys=True, separators=(",", ":"))
        return {
            "output": {"message": {"content": [{"text": text}]}},
            "usage": {"inputTokens": 1, "outputTokens": 1, "totalTokens": 2},
        }


class _RecordingTools:
    """Observe bounded results while delegating every operation to real C09 tools."""

    def __init__(self) -> None:
        self._delegate = AgentTools()
        self.calls: list[str] = []
        self.outputs: dict[str, object] = {}

    def resolve_tool(self, name: str) -> Callable[[object], object]:
        operation = self._delegate.resolve_tool(name)

        def observed(value: object) -> object:
            output = operation(value)
            self.calls.append(name)
            self.outputs[name] = output
            return output

        return observed


class _InjectedS3Client:
    """Minimal deterministic local object store used through the real C14 adapter."""

    def __init__(self, *, false_success: bool = False) -> None:
        self.false_success = false_success
        self.objects: dict[str, dict[str, Any]] = {}
        self.operations: list[str] = []

    def put_object(self, **kwargs: Any) -> dict[str, Any]:
        self.operations.append("put_object")
        if self.false_success:
            return {}
        self.objects[kwargs["Key"]] = dict(kwargs)
        return {"ResponseMetadata": {"HTTPStatusCode": 200}, "ETag": '"c19-local"'}

    def get_object(self, **kwargs: Any) -> dict[str, Any]:
        self.operations.append("get_object")
        value = self.objects[kwargs["Key"]]
        body = value["Body"]
        return {
            "ResponseMetadata": {"HTTPStatusCode": 200},
            "Metadata": dict(value["Metadata"]),
            "ContentLength": len(body),
            "ContentType": value["ContentType"],
            "Body": BytesIO(body),
        }


def _canonical_scenario_bytes(data: Mapping[str, Any]) -> bytes:
    return (json.dumps(data, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()


def load_scenario(path: Path | str = DEFAULT_SCENARIO_PATH) -> LoadedScenario:
    path = Path(path)
    try:
        raw = path.read_bytes()
        data = json.loads(raw)
    except (OSError, json.JSONDecodeError):
        raise ScenarioIntegrityError("scenario_unreadable") from None
    if type(data) is not dict or raw != _canonical_scenario_bytes(data):
        raise ScenarioIntegrityError("scenario_serialization_invalid")
    digest = sha256(raw).hexdigest()
    try:
        parts = path.with_suffix(".sha256").read_text(encoding="utf-8").strip().split()
    except OSError:
        raise ScenarioIntegrityError("scenario_hash_file_missing") from None
    if parts != [digest, path.name]:
        raise ScenarioIntegrityError("scenario_hash_mismatch")
    required = {
        "scenario_schema_version", "scenario_id", "scenario_version", "synthetic",
        "request", "historical_inputs", "provider_trace", "expected",
        "synthetic_approval_actor", "report_claim",
    }
    if set(data) != required:
        raise ScenarioIntegrityError("scenario_fields_invalid")
    if (data["scenario_schema_version"] != SCENARIO_SCHEMA_VERSION
            or data["scenario_version"] != SCENARIO_VERSION
            or data["scenario_id"] != "c19-golden-case-001"
            or data["synthetic"] is not True
            or data["report_claim"] != PASS_CLAIM):
        raise ScenarioIntegrityError("scenario_identity_invalid")
    request = data["request"]
    if type(request) is not dict or set(request) != {"request_id", "quotation_request"}:
        raise ScenarioIntegrityError("scenario_request_invalid")
    quote = request["quotation_request"]
    if type(quote) is not dict or not isinstance(quote.get("items"), list) or not quote["items"]:
        raise ScenarioIntegrityError("scenario_request_invalid")
    for item in quote["items"]:
        if (type(item) is not dict or set(item) != {
                "item_id", "description", "estimated_hours", "hourly_rate"}
                or set(item["estimated_hours"]) != {"value", "unit"}
                or item["estimated_hours"]["unit"] != "hours"
                or set(item["hourly_rate"]) != {"amount", "currency"}
                or item["hourly_rate"]["currency"] != quote.get("currency")):
            raise ScenarioIntegrityError("commercial_semantics_invalid")
    if (not isinstance(data["historical_inputs"], list) or not data["historical_inputs"]
            or any(item.get("synthetic") is not True for item in data["historical_inputs"])
            or not isinstance(data["provider_trace"], list)
            or len(data["provider_trace"]) != len(TOOL_SEQUENCE) + 1):
        raise ScenarioIntegrityError("scenario_evidence_invalid")
    expected = data["expected"]
    if type(expected) is not dict or expected.get("currency") != quote.get("currency"):
        raise ScenarioIntegrityError("expected_values_invalid")
    return LoadedScenario(data=data, sha256=digest, path=path)


def _c17_regression() -> dict[str, Any]:
    with redirect_stdout(StringIO()):
        report = c17.evaluate()
    serialized = c17.serialize_report(report)
    digest = sha256(serialized.encode()).hexdigest()
    return {
        "status": "PASS" if (
            report["overall_status"] == "PASS"
            and report["dataset_version"] == "c17-golden-v1"
            and report["dataset_sha256"] == C17_DATASET_SHA256
            and digest == C17_REPORT_SHA256
        ) else "FAIL",
        "dataset_version": report["dataset_version"],
        "dataset_sha256": report["dataset_sha256"],
        "report_sha256": digest,
        "unchanged": digest == C17_REPORT_SHA256,
    }


def _event_summary(events: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    keys = ("event", "component", "operation", "status")
    counts = Counter(tuple(event.get(key) for key in keys) for event in events)
    return [dict(zip((*keys, "count"), (*values, count), strict=True))
            for values, count in sorted(counts.items())]


def _workbook_summary(blob: bytes, expected: Mapping[str, Any], *, corrupt: bool = False) -> dict[str, Any]:
    if corrupt:
        blob = blob[:-32] + b"c19-corrupt"
    try:
        workbook = load_workbook(BytesIO(blob), data_only=False, keep_links=True)
        quotation = workbook["Quotation"]
        item_count = len(expected["estimated_item_costs"])
        item_costs = [str(quotation.cell(10 + index, 5).value) for index in range(item_count)]
        total = str(quotation.cell(10 + item_count, 5).value)
        formula_free = all(cell.data_type != "f" and cell.hyperlink is None
                           for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
        valid = (
            workbook.sheetnames == ["Quotation", "Risk Analysis", "Historical Evidence"]
            and not workbook._external_links
            and workbook.vba_archive is None
            and item_costs == expected["estimated_item_costs"]
            and total == expected["estimated_total"]
            and quotation["B5"].value == expected["currency"]
            and formula_free
        )
        workbook.close()
        return {
            "status": "PASS" if valid else "FAIL",
            "reloaded": True,
            "reconciled": valid,
            "sheet_names": ["Quotation", "Risk Analysis", "Historical Evidence"],
            "item_costs": item_costs,
            "total": total,
            "currency": expected["currency"] if valid else None,
            "formula_macro_external_link_integrity": valid,
        }
    except Exception:
        return {
            "status": "FAIL", "reloaded": False, "reconciled": False,
            "sheet_names": [], "item_costs": [], "total": None, "currency": None,
            "formula_macro_external_link_integrity": False,
        }


def _blank_details() -> dict[str, Any]:
    return {
        "deterministic_commercial_reconciliation": {},
        "retrieval_comparison_references": {},
        "risk_evidence_references": {},
        "agent_tool_trace_outcome": {},
        "structured_output_validity": {},
        "approval_workflow_result": {},
        "excel_reconciliation_result": {},
        "storage_contract_result": {},
        "observability_contract_result": {},
        "c17_regression_result": {},
    }


def _privacy_check(value: object) -> None:
    if isinstance(value, dict):
        for key, item in value.items():
            if str(key).lower() in _PROHIBITED_KEYS:
                raise ReportIntegrityError("prohibited_report_field")
            _privacy_check(item)
    elif isinstance(value, list):
        for item in value:
            _privacy_check(item)


def validate_report(report: Mapping[str, Any], scenario: LoadedScenario | None = None) -> None:
    if type(report) is not dict or set(report) != _REPORT_KEYS:
        raise ReportIntegrityError("report_fields_invalid")
    _privacy_check(report)
    if report["report_schema_version"] != REPORT_SCHEMA_VERSION:
        raise ReportIntegrityError("report_schema_invalid")
    if report["overall_status"] not in {"PASS", "FAIL", "ERROR"}:
        raise ReportIntegrityError("report_status_invalid")
    if report["synthetic"] is not True or report["limitations"] != ["c13_process_local_state"]:
        raise ReportIntegrityError("claim_boundary_invalid")
    steps = report["required_step_results"]
    if not isinstance(steps, list):
        raise ReportIntegrityError("step_results_invalid")
    ids = [item.get("step_id") for item in steps if type(item) is dict]
    complete = len(ids) == len(REQUIRED_STEPS) and set(ids) == set(REQUIRED_STEPS)
    valid_rows = all(
        type(item) is dict
        and set(item) == {"step_id", "status", "passed", "failure_category"}
        and type(item["passed"]) is bool
        and item["status"] == ("PASS" if item["passed"] else "FAIL")
        and ((item["failure_category"] is None) == item["passed"])
        for item in steps
    )
    if not valid_rows:
        raise ReportIntegrityError("step_result_integrity_invalid")
    if report["overall_status"] != "ERROR" and not complete:
        raise ReportIntegrityError("required_step_accounting_invalid")
    all_pass = complete and all(item["passed"] for item in steps)
    expected_status = "PASS" if all_pass else "FAIL"
    if report["overall_status"] != "ERROR" and report["overall_status"] != expected_status:
        raise ReportIntegrityError("overall_status_misleading")
    if report["overall_status"] == "PASS":
        if report["bounded_failure_category"] is not None or report["claim_boundary"] != PASS_CLAIM:
            raise ReportIntegrityError("pass_claim_invalid")
    elif report["bounded_failure_category"] is None:
        raise ReportIntegrityError("failure_category_missing")
    if scenario is None or report["overall_status"] == "ERROR":
        return
    data = scenario.data
    expected = data["expected"]
    if (report["scenario_id"] != data["scenario_id"]
            or report["scenario_version"] != data["scenario_version"]
            or report["scenario_sha256"] != scenario.sha256):
        raise ReportIntegrityError("scenario_identity_mismatch")
    if report["overall_status"] != "PASS":
        return
    commercial = report["deterministic_commercial_reconciliation"]
    if (commercial.get("expected_total") != expected["estimated_total"]
            or commercial.get("actual_total") != expected["estimated_total"]
            or commercial.get("expected_item_costs") != expected["estimated_item_costs"]
            or commercial.get("actual_item_costs") != expected["estimated_item_costs"]
            or commercial.get("currency") != expected["currency"]):
        raise ReportIntegrityError("commercial_reconciliation_invalid")
    retrieval = report["retrieval_comparison_references"]
    if (retrieval.get("top_similar_quote_ids") != expected["top_similar_quote_ids"]
            or retrieval.get("comparison") != expected["comparison"]):
        raise ReportIntegrityError("retrieval_comparison_invalid")
    risk = report["risk_evidence_references"]
    if (risk.get("evidence_ids") != expected["evidence_ids"]
            or risk.get("selected_evidence_id") != expected["selected_evidence_id"]
            or risk.get("statistics") != expected["risk_statistics"]):
        raise ReportIntegrityError("risk_evidence_invalid")
    trace = report["agent_tool_trace_outcome"]
    if trace.get("tool_sequence") != list(TOOL_SEQUENCE) or trace.get("provider_call_count") != 6:
        raise ReportIntegrityError("agent_trace_invalid")
    if report["structured_output_validity"] != {
            "status": "PASS", "schema_valid": True, "quote_id": expected["quote_id"]}:
        raise ReportIntegrityError("structured_output_invalid")
    approval = report["approval_workflow_result"]
    if approval != {
            "status": "PASS", "review_state": "approved", "preapproval_export_blocked": True,
            "synthetic_test_actor": True, "ai_supplied_approval": False}:
        raise ReportIntegrityError("approval_invalid")
    excel = report["excel_reconciliation_result"]
    if not (excel.get("status") == "PASS" and excel.get("reloaded") is True
            and excel.get("reconciled") is True
            and excel.get("formula_macro_external_link_integrity") is True):
        raise ReportIntegrityError("excel_reconciliation_invalid")
    storage = report["storage_contract_result"]
    if not (storage.get("status") == "PASS" and storage.get("persisted") is True
            and storage.get("retrieved") is True and storage.get("contract_valid") is True
            and storage.get("injected_client") is True and storage.get("real_s3") is False):
        raise ReportIntegrityError("storage_contract_invalid")
    observability = report["observability_contract_result"]
    observed = {(item["event"], item["component"], item["operation"], item["status"])
                for item in observability.get("events", [])}
    required_events = {
        ("api_request", "api", "draft", "success"),
        ("agent_run_finished", "agent", "run", "success"),
        ("bedrock_call", "bedrock", "converse", "success"),
        ("tool_call", "agent_tool", "get_risk_evidence", "success"),
        ("storage_operation", "s3_storage", "persist", "success"),
    }
    if (not required_events <= observed or observability.get("status") != "PASS"
            or observability.get("correlation_present") is not True
            or observability.get("privacy_redaction") is not True
            or observability.get("live_cloudwatch") is not False):
        raise ReportIntegrityError("observability_contract_invalid")
    if report["c17_regression_result"] != {
            "status": "PASS", "dataset_version": "c17-golden-v1",
            "dataset_sha256": C17_DATASET_SHA256, "report_sha256": C17_REPORT_SHA256,
            "unchanged": True}:
        raise ReportIntegrityError("c17_regression_invalid")


def _assemble_report(
    scenario: LoadedScenario,
    ledger: Mapping[str, StepResult],
    details: Mapping[str, Any],
    *,
    force_error: str | None = None,
) -> dict[str, Any]:
    rows = [ledger[name].as_report() for name in REQUIRED_STEPS if name in ledger]
    complete = len(ledger) == len(REQUIRED_STEPS) and set(ledger) == set(REQUIRED_STEPS)
    if force_error is not None or not complete:
        overall = "ERROR"
        failure = force_error or "required_step_accounting_invalid"
    elif all(item.passed for item in ledger.values()):
        overall = "PASS"
        failure = None
    else:
        overall = "FAIL"
        failure = next(item.failure_category for item in ledger.values() if not item.passed)
    report = {
        "report_schema_version": REPORT_SCHEMA_VERSION,
        "scenario_id": scenario.data["scenario_id"],
        "scenario_version": scenario.data["scenario_version"],
        "scenario_sha256": scenario.sha256,
        "synthetic": True,
        "overall_status": overall,
        "required_step_results": rows,
        **details,
        "bounded_failure_category": failure,
        "claim_boundary": PASS_CLAIM if overall == "PASS" else "No Golden Case pass claim is made.",
        "limitations": ["c13_process_local_state"],
    }
    validate_report(report, scenario)
    return report


def _error_report(digest: str, category: str) -> dict[str, Any]:
    details = _blank_details()
    report = {
        "report_schema_version": REPORT_SCHEMA_VERSION,
        "scenario_id": "invalid-scenario",
        "scenario_version": "invalid",
        "scenario_sha256": digest,
        "synthetic": True,
        "overall_status": "ERROR",
        "required_step_results": [],
        **details,
        "bounded_failure_category": category,
        "claim_boundary": "No Golden Case pass claim is made.",
        "limitations": ["c13_process_local_state"],
    }
    validate_report(report)
    return report


def run_once(
    path: Path | str = DEFAULT_SCENARIO_PATH,
    *,
    fault: str | None = None,
) -> dict[str, Any]:
    """Execute one fresh local Golden Case instance and return bounded evidence."""
    if fault is not None and fault not in FAULTS:
        return _error_report("0" * 64, "unknown_fault_injection")
    try:
        scenario = load_scenario(path)
    except ScenarioIntegrityError:
        try:
            digest = sha256(Path(path).read_bytes()).hexdigest()
        except OSError:
            digest = "0" * 64
        return _error_report(digest, "scenario_integrity_error")

    ledger = {name: StepResult(name, False, "not_executed") for name in REQUIRED_STEPS}
    details = _blank_details()

    def mark(name: str, passed: bool, category: str) -> None:
        ledger[name] = StepResult(name, passed, None if passed else category)

    mark("fixture_identity", True, "scenario_integrity_error")
    c17_result = _c17_regression()
    details["c17_regression_result"] = c17_result
    mark("c17_regression", c17_result["status"] == "PASS", "c17_regression_failed")

    provider_client = _ScriptedConverseClient(scenario.data["provider_trace"], fault=fault)
    provider = BedrockConverseAdapter(
        Settings(aws_region="eu-north-1", bedrock_model_id="c19-scripted-model"),
        provider_client,
    )
    tools = _RecordingTools()
    agent = QuotationAgent(provider, tools)  # type: ignore[arg-type]
    store = LocalQuoteStore()
    app = create_app(agent, store=store)
    capture = _EventCapture()
    logger = configure_logging()
    logger.addHandler(capture)
    workbook_blob: bytes | None = None
    try:
        with TestClient(app, raise_server_exceptions=False) as client:
            draft = client.post("/quotes/draft", json=scenario.data["request"])
            if draft.status_code != 200:
                category = "provider_unavailable" if draft.status_code == 503 else "invalid_provider_output"
                mark("fastapi_request_boundary", False, category)
                mark("provider_adapter", False, category)
                mark("quotation_agent", False, category)
                mark("structured_output_validation", False, category)
                details["agent_tool_trace_outcome"] = {
                    "status": "FAIL", "tool_sequence": tools.calls,
                    "provider_call_count": provider_client.call_count,
                    "bounded": True,
                }
                return _assemble_report(scenario, ledger, details)
            mark("fastapi_request_boundary", True, "api_boundary_failed")
            quote_id = draft.json()["quote_id"]
            current, session = store.current(quote_id)
            expected = scenario.data["expected"]

            quote = current.draft_quote.quote
            actual_item_costs = [str(item.estimated_hours.value * item.hourly_rate.amount)
                                 for item in quote.items]
            actual_total = str(quote.estimated_total_cost.amount)
            commercial_pass = (
                quote.quote_id == expected["quote_id"]
                and actual_item_costs == expected["estimated_item_costs"]
                and actual_total == expected["estimated_total"]
                and quote.estimated_total_cost.currency == expected["currency"]
                and str(sum(item.estimated_hours.value for item in quote.items))
                    == expected["total_estimated_hours"]
            )
            details["deterministic_commercial_reconciliation"] = {
                "status": "PASS" if commercial_pass else "FAIL",
                "expected_item_costs": expected["estimated_item_costs"],
                "actual_item_costs": actual_item_costs,
                "expected_total": expected["estimated_total"],
                "actual_total": actual_total,
                "currency": quote.currency,
                "total_estimated_hours": str(sum(item.estimated_hours.value for item in quote.items)),
                "oracle": "independently_authored_fixture",
            }
            mark("core_calculation", commercial_pass, "commercial_reconciliation_failed")

            history = tools.outputs.get("get_historical_quotes")
            similar = tools.outputs.get("find_similar_quotes")
            comparisons = tools.outputs.get("compare_estimate_to_actual")
            statistics = tools.outputs.get("calculate_quote_statistics")
            risk = tools.outputs.get("get_risk_evidence")
            history_pass = (isinstance(history, HistoricalQuotesOutput)
                            and len(history.records) == expected["historical_record_count"])
            mark("historical_retrieval", history_pass, "historical_retrieval_failed")
            similar_ids = ([match.result.quote_id for match in similar.matches]
                           if isinstance(similar, SimilarQuotesOutput) else [])
            similar_pass = similar_ids == expected["top_similar_quote_ids"]
            mark("similar_retrieval", similar_pass, "similar_retrieval_failed")
            comparison = None
            if isinstance(comparisons, ComparisonsOutput):
                comparison = next((item for item in comparisons.comparisons
                                   if item.quote_id == expected["comparison"]["quote_id"]), None)
            actual_comparison = ({
                "quote_id": comparison.quote_id,
                "source_id": comparison.source_id,
                "estimated_hours": str(comparison.hour_variance.estimated),
                "actual_hours": str(comparison.hour_variance.actual),
                "hour_variance": str(comparison.hour_variance.variance),
            } if comparison is not None and comparison.hour_variance is not None else {})
            comparison_pass = actual_comparison == expected["comparison"]
            mark("historical_comparison", comparison_pass, "historical_comparison_failed")
            summary = statistics.summary if isinstance(statistics, StatisticsOutput) else None
            observation_count = len(statistics.sources) if isinstance(statistics, StatisticsOutput) else 0
            actual_stats = {
                "comparable_project_count": observation_count,
                "overrun_count": summary.overrun_count if summary is not None else None,
                "overrun_rate": (str(Decimal(summary.overrun_count) / Decimal(observation_count))
                                 if summary is not None and observation_count else None),
                "average_variance": str(summary.average_variance) if summary is not None else None,
                "median_variance": str(summary.median_variance) if summary is not None else None,
            }
            stats_pass = actual_stats == expected["risk_statistics"]
            mark("deterministic_statistics", stats_pass, "statistics_failed")
            evidence_ids = ([item.evidence_id for item in risk.report.evidence]
                            if isinstance(risk, RiskEvidenceOutput) else [])
            risk_stats = ({
                "comparable_project_count": risk.report.comparable_project_count,
                "overrun_count": risk.report.overrun_count,
                "overrun_rate": str(risk.report.overrun_rate),
                "average_variance": str(risk.report.average_variance),
                "median_variance": str(risk.report.median_variance),
            } if isinstance(risk, RiskEvidenceOutput) else {})
            risk_pass = evidence_ids == expected["evidence_ids"] and risk_stats == expected["risk_statistics"]
            mark("risk_evidence", risk_pass, "risk_evidence_failed")
            details["retrieval_comparison_references"] = {
                "status": "PASS" if similar_pass and comparison_pass else "FAIL",
                "top_similar_quote_ids": similar_ids,
                "comparison": actual_comparison,
            }
            details["risk_evidence_references"] = {
                "status": "PASS" if risk_pass else "FAIL",
                "evidence_ids": evidence_ids,
                "selected_evidence_id": expected["selected_evidence_id"],
                "statistics": risk_stats,
                "provenance_source_count": (
                    len(risk.report.source_quote_ids) if isinstance(risk, RiskEvidenceOutput) else 0),
                "data_origin": "synthetic",
            }
            tool_pass = tools.calls == list(TOOL_SEQUENCE)
            mark("agent_tools", tool_pass, "agent_tool_trace_failed")
            provider_pass = (provider_client.call_count == 6 and provider_client.request_shapes_valid)
            mark("provider_adapter", provider_pass, "provider_adapter_failed")
            agent_pass = current.status is AgentResultStatus.SUCCESS and quote_id == expected["quote_id"]
            mark("quotation_agent", agent_pass, "quotation_agent_failed")
            structured_pass = agent_pass and type(current).__name__ == "AgentResult"
            mark("structured_output_validation", structured_pass, "structured_output_invalid")
            details["agent_tool_trace_outcome"] = {
                "status": "PASS" if tool_pass and provider_pass else "FAIL",
                "tool_sequence": tools.calls,
                "provider_call_count": provider_client.call_count,
                "bounded": True,
            }
            details["structured_output_validity"] = {
                "status": "PASS" if structured_pass else "FAIL",
                "schema_valid": structured_pass,
                "quote_id": quote_id,
            }

            before_approval = client.post(f"/quotes/{quote_id}/export")
            preapproval_blocked = (
                before_approval.status_code == 409
                and before_approval.json() == {"code": "approval_required"}
            )
            if fault == "skip_approval":
                mark("human_approval", False, "approval_skipped")
                details["approval_workflow_result"] = {
                    "status": "FAIL", "review_state": session.state.value,
                    "preapproval_export_blocked": preapproval_blocked,
                    "synthetic_test_actor": True, "ai_supplied_approval": False,
                }
                return _assemble_report(scenario, ledger, details)
            approval = client.post(f"/quotes/{quote_id}/approve", json={
                "status": "approved",
                "reviewer_id": scenario.data["synthetic_approval_actor"],
                "decided_at": "2026-10-05T12:00:00Z",
            })
            approval_pass = approval.status_code == 200 and session.state.value == "approved"
            mark("human_approval", approval_pass and preapproval_blocked, "approval_failed")
            details["approval_workflow_result"] = {
                "status": "PASS" if approval_pass and preapproval_blocked else "FAIL",
                "review_state": session.state.value,
                "preapproval_export_blocked": preapproval_blocked,
                "synthetic_test_actor": True,
                "ai_supplied_approval": False,
            }

            exported = client.post(f"/quotes/{quote_id}/export")
            generation_pass = exported.status_code == 200 and exported.content.startswith(b"PK\x03\x04")
            mark("excel_generation", generation_pass, "excel_generation_failed")
            if generation_pass:
                workbook_blob = exported.content
            excel = (_workbook_summary(workbook_blob, expected,
                                       corrupt=fault == "excel_reconciliation_failure")
                     if workbook_blob is not None else _workbook_summary(b"", expected))
            details["excel_reconciliation_result"] = excel
            mark("excel_reload_reconciliation", excel["status"] == "PASS",
                 "excel_reconciliation_failed")

            if workbook_blob is not None and excel["status"] == "PASS":
                injected_s3 = _InjectedS3Client(false_success=fault == "storage_false_success")
                adapter = S3StorageAdapter(
                    Settings(aws_region="eu-north-1", s3_bucket="aqi-c19-synthetic"), injected_s3)
                artifact = WorkbookArtifact.from_validated_export(
                    request_id=scenario.data["request"]["request_id"],
                    quote_id=quote_id,
                    workbook_bytes=workbook_blob,
                )
                try:
                    stored = adapter.persist(artifact)
                    retrieved = adapter.retrieve(stored.identity)
                    expected_key = derive_object_key(stored.identity)
                    storage_pass = (
                        retrieved.workbook_bytes == workbook_blob
                        and injected_s3.operations == ["put_object", "get_object"]
                        and expected_key in injected_s3.objects
                    )
                except Exception:
                    storage_pass = False
                details["storage_contract_result"] = {
                    "status": "PASS" if storage_pass else "FAIL",
                    "persisted": storage_pass,
                    "retrieved": storage_pass,
                    "contract_valid": storage_pass,
                    "injected_client": True,
                    "real_s3": False,
                    "request_id": scenario.data["request"]["request_id"],
                    "quote_id": quote_id,
                    "object_key_namespace": "generated/v1/",
                }
                mark("storage_adapter", storage_pass, "storage_contract_failed")
            else:
                mark("storage_adapter", False, "storage_not_executed")
    finally:
        logger.removeHandler(capture)

    events = [] if fault == "omit_observability" else _event_summary(capture.events)
    correlation_present = bool(capture.events) and all(
        "request_id" in event for event in capture.events
        if event.get("component") in {"api", "agent", "agent_tool", "bedrock", "s3_storage"}
    )
    observability = {
        "status": "PASS" if events and correlation_present else "FAIL",
        "events": events,
        "correlation_present": correlation_present,
        "privacy_redaction": True,
        "live_cloudwatch": False,
    }
    details["observability_contract_result"] = observability
    mark("observability", observability["status"] == "PASS", "observability_evidence_missing")
    return _assemble_report(scenario, ledger, details)


def run_twice(path: Path | str = DEFAULT_SCENARIO_PATH) -> tuple[dict[str, Any], dict[str, Any], bool]:
    first = run_once(path)
    second = run_once(path)
    return first, second, serialize_report(first) == serialize_report(second)


def serialize_report(
    report: Mapping[str, Any], scenario: LoadedScenario | None = None,
) -> str:
    resolved = scenario
    if resolved is None and report.get("overall_status") != "ERROR":
        resolved = load_scenario()
    validate_report(report, resolved)
    return json.dumps(report, ensure_ascii=False, sort_keys=True, indent=2) + "\n"


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Run the offline V1-C19 Golden Case")
    parser.add_argument("--scenario", type=Path, default=DEFAULT_SCENARIO_PATH)
    args = parser.parse_args(argv)
    with redirect_stdout(StringIO()):
        first, second, equivalent = run_twice(args.scenario)
    report = first if equivalent else _error_report(first.get("scenario_sha256", "0" * 64),
                                                     "two_run_reproducibility_failed")
    print(serialize_report(report), end="")
    return 0 if report["overall_status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())

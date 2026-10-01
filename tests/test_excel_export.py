"""C12 approved export, workbook reconciliation, and adversarial boundaries."""

import json
from datetime import datetime, timezone
from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook

import ai_quotation_intelligence.excel_export as export_module
from ai_quotation_intelligence.agent_tools import AgentTools, RiskEvidenceInput
from ai_quotation_intelligence.bedrock import BedrockResult
from ai_quotation_intelligence.calculation import calculate_item_cost, calculate_quote_total
from ai_quotation_intelligence.data.synthetic_history import generate_synthetic_history
from ai_quotation_intelligence.domain import (
    AgentRequest, AgentResult, AgentResultStatus, ApprovalDecision, ApprovalStatus,
    Money, NewQuoteRequest, QuoteStatus, RiskSeverity, RiskSuggestion,
)
from ai_quotation_intelligence.excel_export import (
    ExportFailure, ExportFailureCode, export_approved_quote,
)
from ai_quotation_intelligence.human_review import ReviewSession
from ai_quotation_intelligence.quotation_agent import QuotationAgent


class ScriptedModel:
    def __init__(self, evidence_id: str) -> None:
        self.outputs = [
            json.dumps({"kind": "tool", "name": "get_risk_evidence", "arguments": {"metric": "hours"}}),
            json.dumps({"kind": "final", "narrative": "Draft for human review.",
                        "evidence_ids": [evidence_id],
                        "risk_suggestions": [{"severity": "medium", "evidence_ids": [evidence_id]}],
                        "missing_information": []}),
        ]
        self.calls = 0

    def invoke(self, prompt: str, *, request_id: str) -> BedrockResult:
        self.calls += 1
        return BedrockResult(AgentResultStatus.SUCCESS, self.outputs.pop(0), request_id)


@pytest.fixture
def approved() -> tuple[ReviewSession, AgentResult, ScriptedModel]:
    historical = generate_synthetic_history()[0].quote
    request = AgentRequest(
        request_id="export-case",
        quotation_request=NewQuoteRequest(
            project_name=historical.project_name,
            currency=historical.currency,
            items=historical.items,
            requested_at=historical.created_at,
        ),
    )
    evidence_id = AgentTools().get_risk_evidence(
        RiskEvidenceInput(request_id="export-evidence")
    ).report.evidence[0].evidence_id
    model = ScriptedModel(evidence_id)
    result = QuotationAgent(model, AgentTools()).run(request)
    assert result.status is AgentResultStatus.SUCCESS
    assert result.draft_quote is not None
    assert len(result.risk_suggestions) == 1
    session = ReviewSession(result)
    session.decide(ApprovalDecision(
        quote_id=result.draft_quote.quote.quote_id,
        status=ApprovalStatus.APPROVED,
        reviewer_id="human-reviewer",
        decided_at=datetime(2026, 9, 30, tzinfo=timezone.utc),
    ), current=result)
    return session, result, model


def failure(code: ExportFailureCode, call) -> None:
    with pytest.raises(ExportFailure) as error:
        call()
    assert error.value.code is code
    assert str(error.value) == code.value


def workbook(blob: bytes):
    return load_workbook(BytesIO(blob), data_only=False, keep_links=True)


def test_approved_export_uses_real_c10_c11_core_and_three_sheets(approved) -> None:
    session, result, model = approved
    before_calls = model.calls
    blob = export_approved_quote(session, current=result)
    assert blob[:2] == b"PK"
    with ZipFile(BytesIO(blob)) as archive:
        assert not any("vbaProject.bin" in name or name.startswith("xl/externalLinks/")
                       for name in archive.namelist())
    book = workbook(blob)
    assert book.sheetnames == ["Quotation", "Risk Analysis", "Historical Evidence"]
    quote = result.draft_quote.quote
    sheet = book["Quotation"]
    assert sheet["B2"].value == quote.quote_id
    assert sheet["B3"].value == result.request_id
    assert sheet["B4"].value == quote.project_name
    assert sheet["B5"].value == quote.currency
    assert sheet["B7"].value == "approved"
    assert sheet["B8"].value == "human-reviewer"
    assert [cell.value for cell in sheet[9]] == [
        "Item ID", "Description", "Estimated hours", "Hourly rate", "Estimated item cost",
    ]
    for row_index, item in enumerate(quote.items, start=10):
        assert sheet.cell(row_index, 1).value == item.item_id
        assert sheet.cell(row_index, 2).value == item.description
        assert Decimal(str(sheet.cell(row_index, 3).value)) == item.estimated_hours.value
        assert Decimal(str(sheet.cell(row_index, 4).value)) == item.hourly_rate.amount
        assert Decimal(str(sheet.cell(row_index, 5).value)) == calculate_item_cost(item).amount
    assert Decimal(str(sheet.cell(sheet.max_row, 5).value)) == calculate_quote_total(quote).amount
    assert "Role" not in [cell.value for cell in sheet[9]]

    risk = result.risk_suggestions[0]
    assert [cell.value for cell in book["Risk Analysis"][4]] == [
        risk.suggestion_id, risk.severity.value, risk.message, risk.evidence_ids[0],
    ]
    assert book["Historical Evidence"]["A4"].value == result.evidence_ids[0]
    assert book["Historical Evidence"]["B4"].value == risk.suggestion_id
    assert "synthetic by default" in book["Historical Evidence"]["A2"].value
    assert book["Historical Evidence"].max_column == 2
    assert model.calls == before_calls  # No C09 discovery or Bedrock call during export.
    assert not hasattr(session, "generate_excel")
    book.close()


def test_semantic_repeatability_without_byte_identity_requirement(approved) -> None:
    session, result, _ = approved
    first = workbook(export_approved_quote(session, current=result))
    second = workbook(export_approved_quote(session, current=result))
    for name in first.sheetnames:
        assert list(first[name].values) == list(second[name].values)
    first.close()
    second.close()


def test_copied_record_and_approved_looking_objects_have_no_export_authority(approved) -> None:
    session, result, _ = approved
    record = session.require_approved(current=result)
    failure(ExportFailureCode.INVALID_INPUT,
            lambda: export_approved_quote(record, current=result))
    failure(ExportFailureCode.INVALID_INPUT,
            lambda: export_approved_quote(record.quote, current=result))
    failure(ExportFailureCode.INVALID_INPUT,
            lambda: export_approved_quote(result.draft_quote, current=result))
    record.quote.status = QuoteStatus.APPROVED
    failure(ExportFailureCode.INVALID_INPUT,
            lambda: export_approved_quote(record, current=result))


def test_unreviewed_rejected_and_unsuccessful_results_never_export(approved) -> None:
    _, result, _ = approved
    pending = ReviewSession(result)
    failure(ExportFailureCode.APPROVAL_REQUIRED,
            lambda: export_approved_quote(pending, current=result))
    rejected = ReviewSession(result)
    rejected.decide(ApprovalDecision(
        quote_id=result.draft_quote.quote.quote_id,
        status=ApprovalStatus.REJECTED, reviewer_id="human-reviewer",
    ), current=result)
    failure(ExportFailureCode.APPROVAL_REQUIRED,
            lambda: export_approved_quote(rejected, current=result))
    for status in (AgentResultStatus.INVALID, AgentResultStatus.UNAVAILABLE,
                   AgentResultStatus.INSUFFICIENT_EVIDENCE):
        bad = AgentResult(request_id=result.request_id, status=status, message="No draft")
        failure(ExportFailureCode.STALE_RESULT,
                lambda bad=bad: export_approved_quote(approved[0], current=bad))


@pytest.mark.parametrize("mutation", [
    "hours", "rate", "total", "quote_id", "request_id", "evidence", "risk", "message",
])
def test_modified_or_substituted_current_result_fails_closed(approved, mutation: str) -> None:
    session, result, _ = approved
    changed = result.model_copy(deep=True)
    if mutation == "hours":
        changed.draft_quote.quote.items[0].estimated_hours.value += Decimal("1")
    elif mutation == "rate":
        changed.draft_quote.quote.items[0].hourly_rate.amount += Decimal("1")
    elif mutation == "total":
        changed.draft_quote.quote.estimated_total_cost.amount += Decimal("1")
    elif mutation == "quote_id":
        changed.draft_quote.quote.quote_id = "draft-forged"
    elif mutation == "request_id":
        changed.request_id = "forged"
    elif mutation == "evidence":
        changed.evidence_ids = ["forged-evidence"]
        changed.draft_quote.evidence_ids = ["forged-evidence"]
    elif mutation == "risk":
        changed.risk_suggestions[0].message = "Different risk"
        changed.draft_quote.risk_suggestions[0].message = "Different risk"
    else:
        changed.message = "Different model narrative"
    failure(ExportFailureCode.STALE_RESULT,
            lambda: export_approved_quote(session, current=changed))


@pytest.mark.parametrize("field", [
    "project", "item", "item_id", "risk", "reviewer", "request", "evidence_id", "suggestion_id",
])
@pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "  ="])
def test_formula_like_untrusted_text_is_inert_after_reload(approved, field: str, prefix: str) -> None:
    _, source, _ = approved
    result = source.model_copy(deep=True)
    text = prefix + "HYPERLINK(\"https://example.invalid\",\"open\")"
    if field == "project":
        result.draft_quote.quote.project_name = text
    elif field == "item":
        result.draft_quote.quote.items[0].description = text
    elif field == "item_id":
        result.draft_quote.quote.items[0].item_id = text
    elif field == "risk":
        suggestion = RiskSuggestion(
            suggestion_id=result.risk_suggestions[0].suggestion_id,
            severity=RiskSeverity.MEDIUM, message=text,
            evidence_ids=result.risk_suggestions[0].evidence_ids,
        )
        result.risk_suggestions = [suggestion]
        result.draft_quote.risk_suggestions = [suggestion]
    elif field == "request":
        result.request_id = text
        result.draft_quote.quote.quote_id = f"draft-{text.strip()}"
    elif field == "evidence_id":
        result.evidence_ids = [text]
        result.draft_quote.evidence_ids = [text]
        result.risk_suggestions[0].evidence_ids = [text]
        result.draft_quote.risk_suggestions[0].evidence_ids = [text]
    elif field == "suggestion_id":
        result.risk_suggestions[0].suggestion_id = text
        result.draft_quote.risk_suggestions[0].suggestion_id = text
    session = ReviewSession(result)
    reviewer = text if field == "reviewer" else "human-reviewer"
    session.decide(ApprovalDecision(
        quote_id=result.draft_quote.quote.quote_id,
        status=ApprovalStatus.APPROVED, reviewer_id=reviewer,
    ), current=result)
    book = workbook(export_approved_quote(session, current=result))
    location = {
        "project": ("Quotation", "B4"),
        "item": ("Quotation", "B10"),
        "item_id": ("Quotation", "A10"),
        "risk": ("Risk Analysis", "C4"),
        "reviewer": ("Quotation", "B8"),
        "request": ("Quotation", "B3"),
        "evidence_id": ("Historical Evidence", "A4"),
        "suggestion_id": ("Risk Analysis", "A4"),
    }[field]
    cell = book[location[0]][location[1]]
    assert cell.value == "'" + text.strip()  # Domain NonEmptyText normalizes outer whitespace.
    assert cell.data_type == "s"
    assert cell.hyperlink is None
    assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)
    book.close()


def test_core_reconciliation_failure_blocks_export(approved, monkeypatch) -> None:
    session, result, _ = approved
    monkeypatch.setattr(export_module, "calculate_quote_total", lambda quote: Money(
        amount=quote.estimated_total_cost.amount + Decimal("1"), currency=quote.currency,
    ))
    failure(ExportFailureCode.COMMERCIAL_MISMATCH,
            lambda: export_approved_quote(session, current=result))


def test_unrepresentable_excel_number_fails_instead_of_silent_rounding(approved) -> None:
    _, source, _ = approved
    result = source.model_copy(deep=True)
    quote = result.draft_quote.quote
    quote.items[0].hourly_rate.amount = Decimal("123456789012345.67")
    quote.estimated_total_cost = calculate_quote_total(
        quote.model_copy(update={"estimated_total_cost": None})
    )
    session = ReviewSession(result)
    session.decide(ApprovalDecision(
        quote_id=quote.quote_id, status=ApprovalStatus.APPROVED,
        reviewer_id="human-reviewer",
    ), current=result)
    failure(ExportFailureCode.WORKBOOK_INVALID,
            lambda: export_approved_quote(session, current=result))


def test_invalid_workbook_and_generation_failure_never_return_partial_success(
    approved, monkeypatch,
) -> None:
    session, result, _ = approved
    monkeypatch.setattr(export_module, "_render", lambda rows: b"not an xlsx")
    failure(ExportFailureCode.WORKBOOK_INVALID,
            lambda: export_approved_quote(session, current=result))
    def broken(rows):
        raise RuntimeError("secret provider payload")
    monkeypatch.setattr(export_module, "_render", broken)
    failure(ExportFailureCode.GENERATION_FAILED,
            lambda: export_approved_quote(session, current=result))


def test_unexpected_row_assembly_failure_is_sanitized(approved, monkeypatch) -> None:
    session, result, _ = approved
    def broken(*args):
        raise RuntimeError("secret internal error")
    monkeypatch.setattr(export_module, "_rows", broken)
    failure(ExportFailureCode.INVALID_INPUT,
            lambda: export_approved_quote(session, current=result))


@pytest.mark.parametrize("mutation", ["formula", "hyperlink", "macro"])
def test_reload_validation_rejects_injected_formula_link_or_macro(
    approved, monkeypatch, mutation: str,
) -> None:
    session, result, _ = approved
    real_render = export_module._render

    def adulterated(rows):
        if mutation == "macro":
            out = BytesIO(real_render(rows))
            with ZipFile(out, mode="a") as archive:
                archive.writestr("xl/vbaProject.bin", b"not a real macro")
            return out.getvalue()
        book = workbook(real_render(rows))
        if mutation == "formula":
            book["Quotation"]["B4"] = "=1+1"
        else:
            book["Quotation"]["B5"].hyperlink = "https://example.invalid"
        out = BytesIO()
        book.save(out)
        book.close()
        return out.getvalue()

    monkeypatch.setattr(export_module, "_render", adulterated)
    failure(ExportFailureCode.WORKBOOK_INVALID,
            lambda: export_approved_quote(session, current=result))


def test_export_does_not_expose_future_actions_or_paths(approved) -> None:
    session, result, _ = approved
    blob = export_approved_quote(session, current=result)
    assert isinstance(blob, bytes)
    assert not hasattr(export_module, "upload")
    assert not hasattr(export_module, "send_email")
    assert not hasattr(export_module, "approve")

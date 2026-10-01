"""C13 HTTP contract with real C10/C11/C12 behavior and isolated model output."""

import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from ai_quotation_intelligence.agent_tools import AgentTools, RiskEvidenceInput
from ai_quotation_intelligence.api import LocalQuoteStore, XLSX_MIME, create_app
from ai_quotation_intelligence.bedrock import BedrockResult
from ai_quotation_intelligence.data.synthetic_history import generate_synthetic_history
from ai_quotation_intelligence.domain import AgentResult, AgentResultStatus
from ai_quotation_intelligence.excel_export import ExportFailure, ExportFailureCode
from ai_quotation_intelligence.human_review import ReviewFailure, ReviewFailureCode, ReviewSession
from ai_quotation_intelligence.quotation_agent import QuotationAgent


def body(request_id: str = "api-case") -> dict:
    quote = generate_synthetic_history()[0].quote
    return {
        "request_id": request_id,
        "quotation_request": {
            "project_name": quote.project_name, "currency": quote.currency,
            "requested_at": quote.created_at.isoformat(),
            "items": [{
                "item_id": item.item_id, "description": item.description,
                "estimated_hours": item.estimated_hours.model_dump(mode="json"),
                "hourly_rate": item.hourly_rate.model_dump(mode="json"),
            } for item in quote.items],
        },
    }


class ScriptedModel:
    def __init__(self) -> None:
        self.prompts: list[str] = []
        self.evidence_id = AgentTools().get_risk_evidence(
            RiskEvidenceInput(request_id="api-evidence")
        ).report.evidence[0].evidence_id

    def invoke(self, prompt: str, *, request_id: str) -> BedrockResult:
        self.prompts.append(prompt)
        payload = (
            {"kind": "tool", "name": "get_risk_evidence", "arguments": {"metric": "hours"}}
            if len(self.prompts) % 2 else
            {"kind": "final", "narrative": "Draft for human review.",
             "evidence_ids": [self.evidence_id], "risk_suggestions": [],
             "missing_information": []}
        )
        return BedrockResult(AgentResultStatus.SUCCESS, json.dumps(payload), request_id)


@pytest.fixture
def live_api() -> tuple[TestClient, LocalQuoteStore, ScriptedModel]:
    model = ScriptedModel()
    store = LocalQuoteStore()
    app = create_app(QuotationAgent(model, AgentTools()), store=store)
    return TestClient(app, raise_server_exceptions=False), store, model


def make_draft(client: TestClient, request_id: str = "api-case") -> str:
    response = client.post("/quotes/draft", json=body(request_id))
    assert response.status_code == 200, response.json()
    return response.json()["quote_id"]


def approve(client: TestClient, quote_id: str, status: str = "approved"):
    return client.post(f"/quotes/{quote_id}/approve", json={
        "status": status, "reviewer_id": "caller-asserted-reviewer",
    })


def test_exact_routes_health_and_openapi(live_api) -> None:
    client, _, model = live_api
    assert client.get("/health").json() == {"status": "ok"}
    assert not model.prompts
    assert client.get("/openapi.json").status_code == 200
    paths = {path for path in client.app.openapi()["paths"]}
    assert paths == {"/health", "/quotes/analyze", "/quotes/draft",
                     "/quotes/{id}", "/quotes/{id}/approve", "/quotes/{id}/export"}
    assert client.app.debug is False
    assert not any(m.cls.__name__ == "CORSMiddleware" for m in client.app.user_middleware)


def test_analyze_is_not_reviewable_and_draft_is(live_api) -> None:
    client, _, _ = live_api
    analysis = client.post("/quotes/analyze", json=body("analysis"))
    assert analysis.status_code == 200
    assert analysis.json()["status"] == "success"
    assert client.get("/quotes/draft-analysis").status_code == 404
    quote_id = make_draft(client)
    view = client.get(f"/quotes/{quote_id}").json()
    assert {key: view[key] for key in ("quote_id", "request_id", "status", "review_state")} == {
        "quote_id": quote_id, "request_id": "api-case", "status": "success",
        "review_state": "awaiting_review"}
    assert view["draft_quote"]["quote"]["quote_id"] == quote_id
    assert view["draft_quote"]["quote"]["estimated_total_cost"] is not None
    assert "reviewer_id" not in view
    assert "review_session" not in view
    assert "review_record" not in view


@pytest.mark.parametrize("mutation", [
    {"status": "approved"}, {"estimated_total_cost": {"amount": "1", "currency": "USD"}},
    {"evidence_ids": ["forged"]}, {"review_record": {"status": "approved"}},
])
def test_top_level_mass_assignment_rejected(live_api, mutation) -> None:
    client, store, _ = live_api
    assert client.post("/quotes/draft", json={**body(), **mutation}).status_code == 422
    assert not store._entries


def test_nested_authority_and_size_rejected(live_api) -> None:
    client, store, _ = live_api
    forged = body()
    forged["quotation_request"]["items"][0]["actual_cost"] = {"amount": "1", "currency": "USD"}
    assert client.post("/quotes/draft", json=forged).status_code == 422
    forged = body()
    forged["quotation_request"]["items"][0]["hourly_rate"]["currency"] = "EUR"
    assert client.post("/quotes/draft", json=forged).status_code == 422
    assert client.post("/quotes/draft", json=body(), headers={"Content-Length": "70000"}).status_code == 413
    huge = body()
    huge["instructions"] = "x" * 70000
    assert client.post("/quotes/draft", json=huge).status_code == 413
    assert not store._entries


def test_approval_export_and_workbook_are_real(live_api) -> None:
    client, _, model = live_api
    quote_id = make_draft(client)
    assert client.post(f"/quotes/{quote_id}/export").status_code == 409
    assert approve(client, quote_id).json() == {
        "quote_id": quote_id, "review_state": "approved", "reviewer_id": "caller-asserted-reviewer"}
    before = len(model.prompts)
    exported = client.post(f"/quotes/{quote_id}/export")
    assert exported.status_code == 200
    assert exported.headers["content-type"] == XLSX_MIME
    assert exported.headers["content-disposition"] == 'attachment; filename="Draft_Quote.xlsx"'
    workbook = load_workbook(BytesIO(exported.content), read_only=True)
    assert workbook.sheetnames == ["Quotation", "Risk Analysis", "Historical Evidence"]
    assert len(model.prompts) == before  # no model call during review/export
    assert approve(client, quote_id).status_code == 409


def test_rejection_and_decision_validation(live_api) -> None:
    client, _, _ = live_api
    quote_id = make_draft(client)
    assert client.post(f"/quotes/{quote_id}/approve", json={"status": "approved"}).status_code == 422
    assert client.post(f"/quotes/{quote_id}/approve", json={"status": "pending", "reviewer_id": "x"}).status_code == 422
    assert client.post(f"/quotes/{quote_id}/approve", json={"status": "approved", "reviewer_id": "x",
                                                             "quote": {"status": "approved"}}).status_code == 422
    assert approve(client, quote_id, "rejected").status_code == 200
    assert client.post(f"/quotes/{quote_id}/export").status_code == 409
    assert approve(client, quote_id).status_code == 409


def test_stale_result_fails_decision_and_export(live_api) -> None:
    client, store, _ = live_api
    quote_id = make_draft(client)
    original, session = store._entries[quote_id]
    changed = original.model_copy(deep=True)
    changed.message = "Changed advisory text"
    store._entries[quote_id] = (changed, session)
    assert approve(client, quote_id).json() == {"code": "stale_review"}
    store._entries[quote_id] = (original, session)
    assert approve(client, quote_id).status_code == 200
    store._entries[quote_id] = (changed, session)
    assert client.post(f"/quotes/{quote_id}/export").json() == {"code": "stale_review"}


def test_store_is_process_local_bounded_and_duplicate_closed(live_api) -> None:
    client, store, _ = live_api
    quote_id = make_draft(client)
    assert client.post("/quotes/draft", json=body()).status_code == 409
    assert TestClient(create_app(QuotationAgent(ScriptedModel(), AgentTools()),
                                 store=LocalQuoteStore())).get(f"/quotes/{quote_id}").status_code == 404
    assert len(store._entries) == 1


def test_session_capacity_fails_closed() -> None:
    model = ScriptedModel()
    client = TestClient(create_app(QuotationAgent(model, AgentTools()),
                                   store=LocalQuoteStore(limit=1)), raise_server_exceptions=False)
    make_draft(client, "one")
    response = client.post("/quotes/draft", json=body("two"))
    assert (response.status_code, response.json()["code"]) == (503, "session_capacity")


def test_invalid_injected_agent_result_cannot_create_review() -> None:
    class ForgedRunner:
        def run(self, value):
            return AgentResult(request_id="wrong", status=AgentResultStatus.SUCCESS)
    response = TestClient(create_app(ForgedRunner()), raise_server_exceptions=False).post(
        "/quotes/draft", json=body())
    assert (response.status_code, response.json()["code"]) == (500, "invalid_agent_result")


def test_concurrent_decisions_only_one_wins(live_api) -> None:
    client, _, _ = live_api
    quote_id = make_draft(client)
    with ThreadPoolExecutor(max_workers=2) as pool:
        codes = list(pool.map(lambda _: approve(client, quote_id).status_code, range(2)))
    assert sorted(codes) == [200, 409]


@pytest.mark.parametrize("status,expected", [
    (AgentResultStatus.INVALID, (422, "agent_invalid")),
    (AgentResultStatus.UNAVAILABLE, (503, "agent_unavailable")),
    (AgentResultStatus.INSUFFICIENT_EVIDENCE, (422, "insufficient_evidence")),
])
def test_c10_failure_mapping(status, expected) -> None:
    class Runner:
        def run(self, value):
            return AgentResult(request_id=value.request_id, status=status)
    response = TestClient(create_app(Runner()), raise_server_exceptions=False).post(
        "/quotes/draft", json=body())
    assert (response.status_code, response.json()["code"]) == expected


def test_provider_and_internal_exceptions_are_sanitized(monkeypatch, live_api) -> None:
    class ExplodingRunner:
        def run(self, value):
            raise RuntimeError("SECRET-BEDROCK-PAYLOAD")
    response = TestClient(create_app(ExplodingRunner()), raise_server_exceptions=False).post(
        "/quotes/analyze", json=body())
    assert response.status_code == 503
    assert "SECRET" not in response.text
    client, _, _ = live_api
    quote_id = make_draft(client)
    assert approve(client, quote_id).status_code == 200
    monkeypatch.setattr("ai_quotation_intelligence.api.export_approved_quote",
                        lambda *a, **k: (_ for _ in ()).throw(RuntimeError("SECRET-PATH")))
    failed = client.post(f"/quotes/{quote_id}/export")
    assert failed.status_code == 500
    assert "SECRET" not in failed.text


def test_export_still_depends_on_c11_gate(monkeypatch, live_api) -> None:
    client, _, _ = live_api
    quote_id = make_draft(client)
    assert approve(client, quote_id).status_code == 200
    def deny(*args, **kwargs):
        raise ReviewFailure(ReviewFailureCode.APPROVAL_REQUIRED)
    monkeypatch.setattr(ReviewSession, "require_approved", deny)
    response = client.post(f"/quotes/{quote_id}/export")
    assert (response.status_code, response.json()["code"]) == (409, "approval_required")
    assert response.headers["content-type"] != XLSX_MIME


@pytest.mark.parametrize("failure,expected", [
    (ExportFailureCode.COMMERCIAL_MISMATCH, (409, "export_reconciliation_failed")),
    (ExportFailureCode.EVIDENCE_MISMATCH, (409, "export_reconciliation_failed")),
    (ExportFailureCode.WORKBOOK_INVALID, (500, "export_failed")),
])
def test_c12_failure_mapping_is_sanitized(monkeypatch, live_api, failure, expected) -> None:
    client, _, _ = live_api
    quote_id = make_draft(client)
    assert approve(client, quote_id).status_code == 200
    def fail(*args, **kwargs):
        raise ExportFailure(failure)
    monkeypatch.setattr("ai_quotation_intelligence.api.export_approved_quote", fail)
    response = client.post(f"/quotes/{quote_id}/export")
    assert (response.status_code, response.json()["code"]) == expected
    assert response.headers["content-type"] != XLSX_MIME


def test_no_future_card_routes_or_headers(live_api) -> None:
    client, _, _ = live_api
    for path in ("/s3", "/deploy", "/quotes/api-case/upload", "/quotes/api-case/email"):
        assert client.post(path).status_code in {404, 405, 413}
    assert client.get("/quotes/not-found").status_code == 404

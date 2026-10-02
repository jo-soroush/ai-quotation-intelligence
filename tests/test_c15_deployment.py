"""Local-only C15 API Gateway/Lambda deployment contract tests."""

import base64
from copy import deepcopy
from io import BytesIO
import json
import os
from pathlib import Path
import re
import subprocess
import sys
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.responses import Response
from mangum import Mangum
from openpyxl import Workbook, load_workbook
import pytest

from ai_quotation_intelligence.api import XLSX_MIME
from deployment import lambda_handler
from scripts.build_c15_lambda import build as build_package
from scripts.c15_signed_health import check_health
from scripts import validate_c15_package as package_validator
from scripts.validate_c15_package import validate as validate_package


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "deployment" / "c15-http-api.json"
LOCK = ROOT / "deployment" / "requirements-lambda.txt"
ROUTES = {
    "GET /health", "POST /quotes/analyze", "POST /quotes/draft",
    "GET /quotes/{id}", "POST /quotes/{id}/approve",
    "POST /quotes/{id}/export",
}


def event(method: str, path: str, *, body: str | None = None, query: str = "") -> dict:
    return {
        "version": "2.0", "routeKey": f"{method} {path}",
        "rawPath": path, "rawQueryString": query,
        "headers": {"host": "example.execute-api.eu-north-1.amazonaws.com",
                    "content-type": "application/json"},
        "requestContext": {
            "accountId": "000000000000", "apiId": "local-test", "domainName": "example.invalid",
            "http": {"method": method, "path": path, "protocol": "HTTP/1.1",
                     "sourceIp": "127.0.0.1", "userAgent": "local-test"},
            "requestId": "local-only", "routeKey": f"{method} {path}",
            "stage": "$default", "time": "02/Oct/2026:00:00:00 +0000", "timeEpoch": 0,
        },
        "body": body, "isBase64Encoded": False,
    }


def test_handler_import_is_local_and_health_is_unchanged(monkeypatch) -> None:
    import boto3

    monkeypatch.setattr(boto3, "client", lambda *a, **k: pytest.fail("AWS client at import/health"))
    result = lambda_handler.handler(event("GET", "/health"), None)
    assert result["statusCode"] == 200
    assert json.loads(result["body"]) == {"status": "ok"}
    assert result["isBase64Encoded"] is False
    assert result["headers"]["content-type"].startswith("application/json")
    assert {route.path for route in lambda_handler.app.routes if route.path in {
        "/health", "/quotes/analyze", "/quotes/draft", "/quotes/{id}",
        "/quotes/{id}/approve", "/quotes/{id}/export",
    }} == {"/health", "/quotes/analyze", "/quotes/draft", "/quotes/{id}",
          "/quotes/{id}/approve", "/quotes/{id}/export"}


def test_fresh_handler_import_does_not_construct_aws_client() -> None:
    script = (
        "import boto3; "
        "boto3.client=lambda *a, **k: (_ for _ in ()).throw(RuntimeError('AWS client forbidden')); "
        "from deployment.lambda_handler import handler; "
        "print(callable(handler))"
    )
    environment = dict(os.environ)
    environment["PYTHONPATH"] = str(ROOT / "src") + os.pathsep + str(ROOT)
    result = subprocess.run([sys.executable, "-c", script], cwd=ROOT,
                            env=environment, capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stderr
    assert result.stdout.strip() == "True"


def test_json_path_and_post_body_mapping() -> None:
    missing = lambda_handler.handler(event("GET", "/quotes/unknown", query="view=1"), None)
    assert missing["statusCode"] == 404
    assert json.loads(missing["body"]) == {"code": "quote_not_found"}
    invalid = lambda_handler.handler(event("POST", "/quotes/analyze", body="{}"), None)
    assert invalid["statusCode"] == 422
    assert json.loads(invalid["body"]) == {"code": "invalid_request"}


def test_binary_xlsx_proxy_response_preserves_bytes_and_headers() -> None:
    workbook = Workbook()
    workbook.active["A1"] = "C15 binary transport"
    stream = BytesIO()
    workbook.save(stream)
    original = stream.getvalue()
    app = FastAPI()

    @app.post("/quotes/demo/export")
    def export() -> Response:
        return Response(original, media_type=XLSX_MIME,
                        headers={"Content-Disposition": 'attachment; filename="Draft_Quote.xlsx"'})

    response = Mangum(app, lifespan="off")(event("POST", "/quotes/demo/export"), None)
    assert response["statusCode"] == 200
    assert response["isBase64Encoded"] is True
    assert base64.b64decode(response["body"], validate=True) == original
    assert response["headers"]["content-type"].startswith(XLSX_MIME)
    assert response["headers"]["content-disposition"] == 'attachment; filename="Draft_Quote.xlsx"'
    assert load_workbook(BytesIO(base64.b64decode(response["body"]))).active["A1"].value == "C15 binary transport"


def test_malformed_event_and_adapter_failure_are_sanitized(monkeypatch) -> None:
    malformed = lambda_handler.handler({"password": "should-not-leak"}, None)
    assert malformed["statusCode"] == 500
    assert "password" not in json.dumps(malformed)
    monkeypatch.setattr(lambda_handler, "_adapter", lambda *_: (_ for _ in ()).throw(RuntimeError("secret-provider-payload")))
    failed = lambda_handler.handler(event("GET", "/health"), None)
    assert failed["statusCode"] == 500
    assert failed["body"] == '{"code":"internal_error"}'
    assert "secret-provider-payload" not in json.dumps(failed)


def assert_template_contract(template: dict) -> None:
    resources = template["Resources"]
    assert template["AWSTemplateFormatVersion"] == "2010-09-09"
    parameters = template["Parameters"]
    assert re.fullmatch(parameters["ArtifactKey"]["AllowedPattern"], f"c15/{'a' * 64}.zip")
    assert not re.fullmatch(parameters["ArtifactKey"]["AllowedPattern"], "../other.zip")
    arn_pattern = parameters["BedrockModelArn"]["AllowedPattern"]
    assert re.fullmatch(arn_pattern, "")
    assert re.fullmatch(arn_pattern, "arn:aws:bedrock:eu-north-1::foundation-model/amazon.nova-micro-v1:0")
    assert re.fullmatch(arn_pattern, "arn:aws:bedrock:eu-north-1:123456789012:inference-profile/eu.amazon.nova-micro-v1")
    assert re.fullmatch(arn_pattern, "arn:aws:bedrock:eu-north-1:123456789012:application-inference-profile/profile-123")
    assert not re.fullmatch(arn_pattern, "*")
    assert not re.fullmatch(arn_pattern, "   ")
    assert not re.fullmatch(arn_pattern, "not-an-arn")
    assert not re.fullmatch(arn_pattern, "arn:aws:bedrock:eu-north-1::inference-profile/profile-123")
    assert not re.fullmatch(arn_pattern, "arn:aws:bedrock:eu-north-1:bad:inference-profile/profile-123")
    assert {value["Properties"]["RouteKey"] for value in resources.values()
            if value["Type"] == "AWS::ApiGatewayV2::Route"} == ROUTES
    assert all(value["Properties"]["AuthorizationType"] == "AWS_IAM"
               for value in resources.values() if value["Type"] == "AWS::ApiGatewayV2::Route")
    assert resources["HttpApi"]["Properties"]["ProtocolType"] == "HTTP"
    assert resources["ApiIntegration"]["Properties"]["PayloadFormatVersion"] == "2.0"
    assert resources["ApiIntegration"]["Properties"]["IntegrationMethod"] == "POST"
    assert resources["ApiIntegration"]["Properties"]["IntegrationUri"]["Fn::Sub"] == (
        "arn:${AWS::Partition}:apigateway:${AWS::Region}:lambda:path/2015-03-31/"
        "functions/${ApiFunction.Arn}/invocations"
    )
    assert resources["ApiFunction"]["Properties"]["Runtime"] == "python3.13"
    assert resources["ApiFunction"]["Properties"]["Handler"] == "lambda_handler.handler"
    assert resources["ApiFunction"]["Properties"]["Architectures"] == ["x86_64"]
    env = resources["ApiFunction"]["Properties"]["Environment"]["Variables"]
    assert "AQI_S3_BUCKET" not in env
    assert not any("key" in field.lower() or "secret" in field.lower() or "token" in field.lower()
                   for field in env)
    policy = resources["LambdaExecutionRole"]["Properties"]["Policies"][0]["PolicyDocument"]
    actions = json.dumps(policy)
    assert "s3:" not in actions and "bedrock:*" not in actions and "logs:*" not in actions
    assert "bedrock:InvokeModel" in actions and "HasBedrockModel" in actions
    bedrock_if = next(statement["Fn::If"] for statement in policy["Statement"] if "Fn::If" in statement)
    assert bedrock_if[0] == "HasBedrockModel"
    bedrock_grant = bedrock_if[1]
    assert bedrock_grant["Action"] == "bedrock:InvokeModel"
    assert bedrock_grant["Resource"] == {"Ref": "BedrockModelArn"}
    assert bedrock_grant["Resource"] != "*"
    assert "logs:CreateLogStream" in actions and "logs:PutLogEvents" in actions
    assert not any(value["Type"] in {"AWS::S3::Bucket", "AWS::DynamoDB::Table"}
                   for value in resources.values())
    assert "CorsConfiguration" not in resources["HttpApi"]["Properties"]


def test_template_routes_access_iam_and_runtime_permissions() -> None:
    assert_template_contract(json.loads(TEMPLATE.read_text(encoding="utf-8")))


@pytest.mark.parametrize("mutation", ["remove_route", "anonymous_route", "wrong_target",
                                      "s3_grant", "wildcard_cors", "wildcard_model_arn"])
def test_targeted_template_mutations_are_detected(mutation: str) -> None:
    template = json.loads(TEMPLATE.read_text(encoding="utf-8"))
    altered = deepcopy(template)
    resources = altered["Resources"]
    if mutation == "remove_route":
        del resources["RouteExport"]
    elif mutation == "anonymous_route":
        resources["RouteExport"]["Properties"]["AuthorizationType"] = "NONE"
    elif mutation == "wrong_target":
        resources["ApiIntegration"]["Properties"]["PayloadFormatVersion"] = "1.0"
    elif mutation == "s3_grant":
        resources["LambdaExecutionRole"]["Properties"]["Policies"][0]["PolicyDocument"]["Statement"].append({"Action": "s3:*"})
    elif mutation == "wildcard_cors":
        resources["HttpApi"]["Properties"]["CorsConfiguration"] = {"AllowOrigins": ["*"]}
    else:
        altered["Parameters"]["BedrockModelArn"]["AllowedPattern"] = ".*"
    with pytest.raises(AssertionError):
        assert_template_contract(altered)


def test_lock_is_pinned_and_package_builder_never_invokes_aws() -> None:
    lines = [line for line in LOCK.read_text().splitlines() if line and not line.startswith("#")]
    assert all(re.fullmatch(r"[A-Za-z0-9_-]+==[A-Za-z0-9_.]+", line) for line in lines)
    assert len(lines) == len({line.split("==")[0].lower().replace("_", "-") for line in lines})
    assert {"mangum==0.20.0", "openpyxl==3.1.5", "boto3==1.43.96",
            "botocore==1.43.96", "opentelemetry-api==1.45.0"} <= set(lines)
    builder = (ROOT / "scripts" / "build_c15_lambda.py").read_text()
    assert "aws " not in builder and "boto3" not in builder
    assert "manylinux2014_x86_64" in builder


def test_builder_refuses_overwrite_before_pip(tmp_path: Path) -> None:
    artifact = tmp_path / "already-exists.zip"
    artifact.write_bytes(b"do-not-overwrite")
    with pytest.raises(ValueError, match="already exists"):
        build_package(artifact)
    assert artifact.read_bytes() == b"do-not-overwrite"


@pytest.mark.parametrize(("api_id", "region"), [
    ("other.example.com", "eu-north-1"),
    ("abcdefghij", "eu-north-1/attacker"),
    ("abcdefghij\r\nHost:evil", "eu-north-1"),
])
def test_future_signed_smoke_refuses_untrusted_destination(monkeypatch, api_id: str, region: str) -> None:
    import boto3

    monkeypatch.setattr(boto3, "Session", lambda **_: pytest.fail("credentials resolved for invalid target"))
    with pytest.raises(ValueError):
        check_health(api_id, region)


def test_no_static_credentials_or_public_cors_in_deployment_files() -> None:
    files = [TEMPLATE, ROOT / "deployment" / "lambda_handler.py",
             ROOT / "scripts" / "build_c15_lambda.py"]
    joined = "\n".join(path.read_text(encoding="utf-8") for path in files)
    assert not re.search(r"\b(?:AKIA|ASIA)[A-Z0-9]{16}\b", joined)
    assert "aws_secret_access_key" not in joined
    assert "Access-Control-Allow-Origin" not in joined
    assert "public-read" not in joined


def test_package_validator_rejects_incomplete_or_sensitive_archives(tmp_path: Path) -> None:
    artifact = tmp_path / "invalid.zip"
    with ZipFile(artifact, "w") as archive:
        archive.writestr("lambda_handler.py", "")
        archive.writestr(".aws/credentials", "fake")
    with pytest.raises(ValueError, match="missing application"):
        validate_package(artifact)


def test_runtime_metadata_closure_rejects_original_missing_dependency(tmp_path: Path, monkeypatch) -> None:
    lock = tmp_path / "requirements-lambda.txt"
    lock.write_text("fastapi==0.142.2\n", encoding="utf-8")
    monkeypatch.setattr(package_validator, "LOCK", lock)
    archive_path = tmp_path / "incomplete.zip"
    with ZipFile(archive_path, "w") as archive:
        archive.writestr("fastapi-0.142.2.dist-info/METADATA",
                         "Name: fastapi\nVersion: 0.142.2\nRequires-Dist: opentelemetry-api>=1.44.0\n")
    with ZipFile(archive_path) as archive, pytest.raises(
        ValueError, match="incomplete or incompatible Lambda runtime dependency closure"
    ):
        package_validator._validate_dependency_closure(archive)


def test_runtime_metadata_closure_accepts_complete_transitive_chain(tmp_path: Path, monkeypatch) -> None:
    lock = tmp_path / "requirements-lambda.txt"
    lock.write_text("fastapi==0.142.2\nopentelemetry-api==1.45.0\n"
                    "typing-extensions==4.16.0\n", encoding="utf-8")
    monkeypatch.setattr(package_validator, "LOCK", lock)
    archive_path = tmp_path / "complete.zip"
    with ZipFile(archive_path, "w") as archive:
        archive.writestr("fastapi-0.142.2.dist-info/METADATA",
                         "Name: fastapi\nVersion: 0.142.2\nRequires-Dist: opentelemetry-api>=1.44.0\n")
        archive.writestr("opentelemetry_api-1.45.0.dist-info/METADATA",
                         "Name: opentelemetry-api\nVersion: 1.45.0\n"
                         "Requires-Dist: typing-extensions>=4.5.0\n")
        archive.writestr("typing_extensions-4.16.0.dist-info/METADATA",
                         "Name: typing_extensions\nVersion: 4.16.0\n")
    with ZipFile(archive_path) as archive:
        package_validator._validate_dependency_closure(archive)

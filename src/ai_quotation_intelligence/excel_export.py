"""C12 in-memory Excel rendering of a current C11-approved quotation.

The held ReviewSession, not copied review metadata, grants export eligibility.
Core supplies all commercial arithmetic; this module only renders and checks it.
"""

from decimal import Decimal
from enum import StrEnum
from io import BytesIO
from zipfile import ZipFile

from ai_quotation_intelligence.calculation import calculate_item_cost, calculate_quote_total
from ai_quotation_intelligence.domain import AgentResult, ApprovalStatus, QuoteStatus
from ai_quotation_intelligence.human_review import (
    ReviewFailure,
    ReviewFailureCode,
    ReviewRecord,
    ReviewSession,
)


SHEETS = ("Quotation", "Risk Analysis", "Historical Evidence")
_HISTORY_NOTE = (
    "Approved evidence references only; historical metrics and source records are not "
    "bound to this result. V1 portfolio history is synthetic by default."
)


class ExportFailureCode(StrEnum):
    INVALID_INPUT = "invalid_input"
    APPROVAL_REQUIRED = "approval_required"
    STALE_RESULT = "stale_result"
    COMMERCIAL_MISMATCH = "commercial_mismatch"
    EVIDENCE_MISMATCH = "evidence_mismatch"
    WORKBOOK_INVALID = "workbook_invalid"
    GENERATION_FAILED = "generation_failed"
    DEPENDENCY_UNAVAILABLE = "dependency_unavailable"


class ExportFailure(Exception):
    """Typed, non-sensitive export failure."""

    def __init__(self, code: ExportFailureCode) -> None:
        self.code = code
        super().__init__(code.value)


def _safe_text(value: str) -> str:
    """Keep untrusted text inert even when a spreadsheet trims leading space."""
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _approved_record(session: object, current: object) -> ReviewRecord:
    if type(session) is not ReviewSession:
        raise ExportFailure(ExportFailureCode.INVALID_INPUT)
    try:
        record = session.require_approved(current=current)
    except ReviewFailure as exc:
        code = (ExportFailureCode.APPROVAL_REQUIRED
                if exc.code is ReviewFailureCode.APPROVAL_REQUIRED
                else ExportFailureCode.STALE_RESULT)
        raise ExportFailure(code) from None
    except Exception:
        raise ExportFailure(ExportFailureCode.INVALID_INPUT) from None
    try:
        clean = ReviewRecord.model_validate_json(record.model_dump_json())
    except Exception:
        raise ExportFailure(ExportFailureCode.INVALID_INPUT) from None
    if (clean.decision.status is not ApprovalStatus.APPROVED
            or clean.quote.status is not QuoteStatus.APPROVED):
        raise ExportFailure(ExportFailureCode.APPROVAL_REQUIRED)
    return clean


def _commercial_truth(record: ReviewRecord) -> tuple[tuple[Decimal, ...], Decimal]:
    try:
        quote = record.quote
        costs = tuple(calculate_item_cost(item) for item in quote.items)
        total = calculate_quote_total(quote)
        if (quote.estimated_total_cost != total
                or record.reviewed_draft.quote.estimated_total_cost != total
                or any(cost.currency != quote.currency for cost in costs)):
            raise ValueError("commercial mismatch")
        return tuple(cost.amount for cost in costs), total.amount
    except Exception:
        raise ExportFailure(ExportFailureCode.COMMERCIAL_MISMATCH) from None


def _evidence_links(record: ReviewRecord) -> tuple[tuple[str, str], ...]:
    try:
        evidence_ids = tuple(record.evidence_ids)
        suggestions = record.reviewed_draft.risk_suggestions
        if (not evidence_ids or len(evidence_ids) != len(set(evidence_ids))
                or evidence_ids != tuple(record.reviewed_draft.evidence_ids)
                or len({item.suggestion_id for item in suggestions}) != len(suggestions)):
            raise ValueError("invalid evidence identities")
        links: list[tuple[str, str]] = []
        for evidence_id in evidence_ids:
            related = [item.suggestion_id for item in suggestions
                       if evidence_id in item.evidence_ids]
            links.extend((evidence_id, suggestion_id) for suggestion_id in related)
            if not related:
                links.append((evidence_id, ""))
        if any(not set(item.evidence_ids) <= set(evidence_ids)
               or len(item.evidence_ids) != len(set(item.evidence_ids))
               for item in suggestions):
            raise ValueError("unbound risk evidence")
        return tuple(links)
    except Exception:
        raise ExportFailure(ExportFailureCode.EVIDENCE_MISMATCH) from None


def _rows(record: ReviewRecord, costs: tuple[Decimal, ...], total: Decimal
          ) -> dict[str, list[tuple[object, ...]]]:
    quote = record.quote
    quotation: list[tuple[object, ...]] = [
        ("Quotation",),
        ("Quote ID", quote.quote_id),
        ("Request ID", record.request_id),
        ("Project", quote.project_name),
        ("Currency", quote.currency),
        ("Quote data origin", quote.data_origin.value),
        ("Review state", "approved"),
        ("Reviewer reference (caller-asserted)", record.decision.reviewer_id),
        ("Item ID", "Description", "Estimated hours", "Hourly rate", "Estimated item cost"),
    ]
    quotation.extend(
        (item.item_id, item.description, item.estimated_hours.value,
         item.hourly_rate.amount, cost)
        for item, cost in zip(quote.items, costs, strict=True)
    )
    quotation.append(("TOTAL", "", "", "", total))

    risks: list[tuple[object, ...]] = [
        ("Risk Analysis",),
        ("Only suggestions bound to the approved result",),
        ("Suggestion ID", "Severity", "Suggestion", "Evidence ID"),
    ]
    for suggestion in record.reviewed_draft.risk_suggestions:
        risks.extend(
            (suggestion.suggestion_id, suggestion.severity.value,
             suggestion.message, evidence_id)
            for evidence_id in suggestion.evidence_ids
        )

    history: list[tuple[object, ...]] = [
        ("Historical Evidence",),
        (_HISTORY_NOTE,),
        ("Approved evidence ID", "Linked approved suggestion ID"),
    ]
    history.extend(_evidence_links(record))
    return dict(zip(SHEETS, (quotation, risks, history), strict=True))


def _openpyxl():
    try:
        from openpyxl import Workbook, load_workbook
    except ImportError:
        raise ExportFailure(ExportFailureCode.DEPENDENCY_UNAVAILABLE) from None
    return Workbook, load_workbook


def _render(rows: dict[str, list[tuple[object, ...]]]) -> bytes:
    Workbook, _ = _openpyxl()
    from openpyxl.styles import Font, PatternFill
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SHEETS:
        sheet = workbook.create_sheet(name)
        for row_index, row in enumerate(rows[name], start=1):
            for column_index, value in enumerate(row, start=1):
                cell = sheet.cell(row_index, column_index)
                if isinstance(value, str):
                    cell.value = _safe_text(value)
                    cell.data_type = "s"
                    cell.number_format = "@"
                else:
                    cell.value = value
                    cell.number_format = "General"  # Never visually round a valid Core value to fixed decimals.
        header_row = 9 if name == "Quotation" else 3
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DCEAF5")
        sheet["A1"].font = Font(bold=True, size=14)
        sheet.freeze_panes = f"A{header_row + 1}"
        for column, width in {"A": 34, "B": 52, "C": 24, "D": 25, "E": 26}.items():
            sheet.column_dimensions[column].width = width
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _same_cell(actual: object, expected: object) -> bool:
    if isinstance(expected, Decimal):
        if isinstance(actual, bool) or not isinstance(actual, (int, float)):
            return False
        return Decimal(str(actual)) == expected
    if isinstance(expected, str):
        if expected == "" and actual is None:
            return True  # openpyxl omits empty-string cells when saving.
        return actual == _safe_text(expected)
    return actual == expected


def _validate(blob: bytes, rows: dict[str, list[tuple[object, ...]]]) -> None:
    try:
        _, load_workbook = _openpyxl()
        with ZipFile(BytesIO(blob)) as archive:
            names = archive.namelist()
            if (any("vbaProject.bin" in name or name.startswith("xl/externalLinks/")
                    for name in names)):
                raise ValueError("macro or external link")
        workbook = load_workbook(BytesIO(blob), data_only=False, keep_links=True)
        if (workbook.sheetnames != list(SHEETS)
                or workbook._external_links
                or workbook.vba_archive is not None):
            raise ValueError("invalid workbook structure")
        for name in SHEETS:
            sheet = workbook[name]
            expected_rows = rows[name]
            if sheet.sheet_state != "visible" or sheet.max_row != len(expected_rows):
                raise ValueError("unexpected sheet rows")
            for row_index, expected_row in enumerate(expected_rows, start=1):
                for column_index in range(1, max(sheet.max_column, len(expected_row)) + 1):
                    cell = sheet.cell(row_index, column_index)
                    expected = (expected_row[column_index - 1]
                                if column_index <= len(expected_row) else None)
                    if (cell.data_type == "f" or cell.hyperlink is not None
                            or not _same_cell(cell.value, expected)):
                        raise ValueError("unexpected workbook cell")
        workbook.close()
    except ExportFailure:
        raise
    except Exception:
        raise ExportFailure(ExportFailureCode.WORKBOOK_INVALID) from None


def export_approved_quote(session: ReviewSession, *, current: AgentResult) -> bytes:
    """Return a validated .xlsx; never persist or infer approval."""
    record = _approved_record(session, current)
    costs, total = _commercial_truth(record)
    try:
        rows = _rows(record, costs, total)
    except ExportFailure:
        raise
    except Exception:
        raise ExportFailure(ExportFailureCode.INVALID_INPUT) from None
    try:
        blob = _render(rows)
    except ExportFailure:
        raise
    except Exception:
        raise ExportFailure(ExportFailureCode.GENERATION_FAILED) from None
    _validate(blob, rows)
    return blob


__all__ = ["ExportFailure", "ExportFailureCode", "export_approved_quote"]
